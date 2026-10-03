"""Committed Trainer / Answer Key presentation pair."""

from __future__ import annotations

import shutil

from openpyxl import load_workbook

from composer.tests.test_learner_ready_presentation import (
    ROOT,
    WHITE_RGBS,
    _assert_answer_key_no_yellow,
    _assert_fresh_visible_style,
    _build_canonical,
    _fill_rgb,
)
from legacy.trainer.checker import check_workbook
from modeler.build_bav import JUDGMENT_RESPONSE_COLS, _judgment_case_rows
from modeler.semantic_io import group_components_by_family, load_semantic_map, parse_cell_ref
from modeler.workbook import JUDGMENT_SHEET, NORMALIZATION_JUDGMENT_SHEET

def test_committed_canonical_pair_matches_minimal_style_contract(tmp_path):
    committed_trainer = ROOT / "legacy" / "example" / "DEMO_HK_Trainer.xlsx"
    committed_answer = ROOT / "legacy" / "example" / "DEMO_HK_Answer_Key.xlsx"
    committed_smap = load_semantic_map(committed_answer)
    assert len(group_components_by_family(committed_smap)) == 78
    assert len(committed_smap.all_ordered()) == 332
    for suffix in (".component_map.json", ".trainer.json", ".assumptions.json"):
        assert not committed_trainer.with_suffix(suffix).exists()

    fresh_trainer, fresh_answer = _build_canonical(tmp_path)
    fresh_smap = load_semantic_map(fresh_answer)
    assert [comp.id for comp in fresh_smap.all_ordered()] == [
        comp.id for comp in committed_smap.all_ordered()
    ]
    practice = {(comp.tab, comp.cell) for comp in fresh_smap.all_ordered()}
    _assert_fresh_visible_style(fresh_trainer, practice_cells=practice, role="trainer")
    _assert_fresh_visible_style(fresh_answer, practice_cells=practice, role="answer_key")
    _assert_answer_key_no_yellow(fresh_answer)

    check_dir = tmp_path / "canonical_check_copy"
    check_dir.mkdir()
    check_trainer = check_dir / fresh_trainer.name
    shutil.copy2(fresh_trainer, check_trainer)
    shutil.copy2(fresh_answer, check_dir / fresh_answer.name)
    for sidecar in (
        fresh_answer.with_suffix(".component_map.json"),
        fresh_answer.with_suffix(".assumptions.json"),
        fresh_answer.with_suffix(".trainer.json"),
    ):
        if sidecar.is_file():
            shutil.copy2(sidecar, check_dir / sidecar.name)
    summary = check_workbook(check_trainer)
    assert (summary.correct, summary.incorrect, summary.blank) == (0, 0, 332)


def test_saved_reopened_pair_answer_key_white_with_both_judgment_modules(tmp_path):
    trainer, answer = _build_canonical(tmp_path)
    reopened = tmp_path / "reopened"
    reopened.mkdir()
    trainer_r = reopened / trainer.name
    answer_r = reopened / answer.name
    shutil.copy2(trainer, trainer_r)
    shutil.copy2(answer, answer_r)
    for sidecar in (
        answer.with_suffix(".component_map.json"),
        answer.with_suffix(".assumptions.json"),
        answer.with_suffix(".trainer.json"),
    ):
        if sidecar.is_file():
            shutil.copy2(sidecar, reopened / sidecar.name)
    for path in (trainer_r, answer_r):
        wb = load_workbook(path, data_only=False)
        wb.save(path)
        wb.close()

    smap = load_semantic_map(answer_r)
    practice = {(comp.tab, comp.cell) for comp in smap.all_ordered()}
    _assert_fresh_visible_style(trainer_r, practice_cells=practice, role="trainer")
    _assert_fresh_visible_style(answer_r, practice_cells=practice, role="answer_key")
    _assert_answer_key_no_yellow(answer_r)

    trainer_wb = load_workbook(trainer_r, data_only=False)
    answer_wb = load_workbook(answer_r, data_only=False)
    visible_t = [ws.title for ws in trainer_wb.worksheets if ws.sheet_state == "visible"]
    visible_a = [ws.title for ws in answer_wb.worksheets if ws.sheet_state == "visible"]
    assert "Overview" in visible_a
    assert "Trainer" not in visible_a
    assert "Trainer" in visible_t
    assert [title for title in visible_t if title != "Trainer"] == visible_a
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        trainer_cell = trainer_wb[comp.tab].cell(row=row, column=col)
        answer_cell = answer_wb[comp.tab].cell(row=row, column=col)
        assert trainer_cell.value is None
        assert trainer_cell.comment is None
        assert _fill_rgb(trainer_cell) == "FFFF00"
        assert isinstance(answer_cell.value, str) and answer_cell.value.startswith("=")
        assert answer_cell.value == comp.formula
        assert answer_cell.comment is not None
        assert (answer_cell.comment.text or "").strip()
        assert _fill_rgb(answer_cell) in WHITE_RGBS
        assert trainer_cell.font.name == answer_cell.font.name == "Aptos Narrow"
        assert trainer_cell.font.size == answer_cell.font.size == 11
        assert trainer_cell.font.bold is False
        assert answer_cell.font.bold is False

    for sheet in (JUDGMENT_SHEET, NORMALIZATION_JUDGMENT_SHEET):
        assert sheet in trainer_wb.sheetnames
        assert sheet in answer_wb.sheetnames
        rows = list(_judgment_case_rows(answer_wb[sheet]))
        assert rows, f"expected non-empty cases on {sheet}"
        for row in rows:
            for col in JUDGMENT_RESPONSE_COLS:
                answer_cell = answer_wb[sheet].cell(row, col)
                trainer_cell = trainer_wb[sheet].cell(row, col)
                assert answer_cell.value not in (None, "")
                assert _fill_rgb(answer_cell) in WHITE_RGBS
                assert trainer_cell.value is None
                assert trainer_cell.comment is None
                assert _fill_rgb(trainer_cell) == "FFFF00"
    trainer_wb.close()
    answer_wb.close()

