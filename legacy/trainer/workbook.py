"""Legacy Trainer derivation — practice blanking, chrome and Trainer-only helpers."""

from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from modeler.build_bav import (
    JUDGMENT_RESPONSE_COLS,
    _apply_minimal_style,
    _financials_for_opening,
    _judgment_case_rows,
    finalize_bav,
)
from modeler.check_context import CHECK_CONTEXT_SHEET
from modeler.engine.component_catalog import is_operating_kpi_source_identity
from modeler.engine.map_embed import COMPONENT_MAP_SHEET
from modeler.engine.semantic_map import SemanticMap
from modeler.semantic_io import group_components_by_family, load_semantic_map, parse_cell_ref
from modeler.workbook import JUDGMENT_SHEET, NORMALIZATION_JUDGMENT_SHEET

PRACTICE_FILL = PatternFill("solid", start_color="FFFF00")

TRAINER_INDEX_INSTRUCTION = (
    "Complete each historical model-construction and earnings-quality diagnostic "
    "formula schedule left-to-right in dependency order. "
    "Run Check to validate the yellow formula cells against the treatment currently "
    "selected in Accounting Judgment column F (blank F uses the supplied reference "
    "treatment). Generated Condensed Financials classification links are system-"
    "controlled and validated by Check—do not edit them or columns D:E. When "
    "Normalization Judgment is present, choose Recurring vs Non-recurring in column F "
    "(blank uses the supplied reference); Earnings Normalization formulas are checked "
    "against that choice. Also complete Accounting Judgment and Normalization Judgment "
    "rationale/consequence when cases are present; those responses are not graded by "
    "Check. Compare them with the matching BAV."
)

_TRAINER_SIDECAR_SUFFIXES = (
    ".component_map.json",
    ".trainer.json",
    ".assumptions.json",
)


def _is_primary_bav_path(path: Path) -> bool:
    stem = Path(path).stem
    return stem.endswith("_BAV") and not stem.endswith("_BAV_Trainer")


def remove_trainer_sidecars(trainer_path: Path) -> None:
    """Delete stale Trainer-only answer-bearing sidecars (idempotent)."""
    trainer_path = Path(trainer_path)
    if _is_primary_bav_path(trainer_path):
        return
    for suffix in _TRAINER_SIDECAR_SUFFIXES:
        trainer_path.with_suffix(suffix).unlink(missing_ok=True)


class TrainingWorkbookGenerator:
    """Optionally derive a sanitized Trainer from a completed professional BAV."""

    def __init__(
        self,
        answer_key_path: Path,
        semantic_map: SemanticMap | None = None,
        financials=None,
    ):
        self.answer_key_path = Path(answer_key_path)
        self.bav_path = self.answer_key_path
        self.semantic_map = semantic_map or load_semantic_map(self.bav_path)
        self.financials = financials

    def finalize_bav(self) -> Path:
        """Compatibility: finalize through Modeler. Does not implement BAV construction."""
        if self.financials is None:
            self.financials = _financials_for_opening(self.bav_path)
        return finalize_bav(self.bav_path, self.semantic_map, self.financials)

    def derive_trainer(self, trainer_path: Path) -> Path:
        """Copy the finalized BAV and sanitize a Trainer without mutating the BAV."""
        trainer_path = Path(trainer_path)
        shutil.copy2(self.bav_path, trainer_path)
        wb = load_workbook(trainer_path)
        self._add_trainer_ui(wb)
        _apply_minimal_style(wb)
        self._blank_trainer_practice_cells(wb)
        self._blank_trainer_judgment_cells(wb)
        self._blank_trainer_normalization_judgment_cells(wb)
        self._sanitize_trainer_answer_stores(wb)
        wb.save(trainer_path)
        wb.close()
        remove_trainer_sidecars(trainer_path)
        return trainer_path

    def generate(self, trainer_path: Path) -> tuple[Path, Path]:
        """Finalize the BAV, then derive a sanitized Trainer from that completed model."""
        self.finalize_bav()
        self.derive_trainer(trainer_path)
        return trainer_path, self.bav_path

    def _blank_trainer_practice_cells(self, wb) -> None:
        for comp in self.semantic_map.all_ordered():
            if is_operating_kpi_source_identity(comp):
                continue
            if comp.tab not in wb.sheetnames:
                continue
            ws = wb[comp.tab]
            row, col = parse_cell_ref(comp.cell)
            cell = ws.cell(row=row, column=col)
            cell.value = None
            cell.fill = PRACTICE_FILL
            cell.comment = None

    def _blank_trainer_judgment_cells(self, wb) -> None:
        if JUDGMENT_SHEET not in wb.sheetnames:
            return
        ws = wb[JUDGMENT_SHEET]
        for row in _judgment_case_rows(ws):
            for col in JUDGMENT_RESPONSE_COLS:
                cell = ws.cell(row=row, column=col)
                cell.value = None
                cell.fill = PRACTICE_FILL
                cell.comment = None

    def _blank_trainer_normalization_judgment_cells(self, wb) -> None:
        if NORMALIZATION_JUDGMENT_SHEET not in wb.sheetnames:
            return
        ws = wb[NORMALIZATION_JUDGMENT_SHEET]
        for row in _judgment_case_rows(ws):
            for col in JUDGMENT_RESPONSE_COLS:
                cell = ws.cell(row=row, column=col)
                cell.value = None
                cell.fill = PRACTICE_FILL
                cell.comment = None

    def _sanitize_trainer_answer_stores(self, wb) -> None:
        """Remove answer-bearing hidden sheets from the Trainer only."""
        for name in (
            COMPONENT_MAP_SHEET,
            CHECK_CONTEXT_SHEET,
            "_RefFormulas",
            "_RefValues",
            "_TrainerMeta",
        ):
            if name in wb.sheetnames:
                del wb[name]

    def _add_trainer_ui(self, wb) -> None:
        if "Trainer" in wb.sheetnames:
            del wb["Trainer"]
        ws = wb.create_sheet("Trainer", 0)
        ws.sheet_view.showGridLines = False
        ws["A1"] = "BAV Excel Trainer"
        ws["A2"] = TRAINER_INDEX_INSTRUCTION
        if "Build Status" in wb.sheetnames:
            ws["G1"] = "Current Progress"
            ws["G1"].hyperlink = "#'Build Status'!A1"
        headers = ["Order", "Schedule", "Period scope", "Tab", "Practice cells", "Depends on"]
        for j, h in enumerate(headers, start=1):
            ws.cell(row=4, column=j, value=h)
        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 36
        ws.column_dimensions["C"].width = 22
        ws.column_dimensions["D"].width = 22
        ws.column_dimensions["E"].width = 28
        ws.column_dimensions["F"].width = 28

        for i, group in enumerate(group_components_by_family(self.semantic_map), start=5):
            ws.cell(row=i, column=1, value=group["family_order"])
            ws.cell(row=i, column=2, value=group["title"])
            ws.cell(row=i, column=3, value=group["period_scope"])
            ws.cell(row=i, column=4, value=group["tab"])
            ws.cell(row=i, column=5, value=group["practice_cells"])
            ws.cell(row=i, column=6, value=group["depends_on"])
