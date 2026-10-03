"""Disclosure-led historical revenue-driver tests and workbook presentation."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path
import shutil

import pytest
from openpyxl import load_workbook

from extractor.data.historical_strategy import (
    ROLE_OBJECTIVE,
    ROLE_OPERATING_USE,
    ROLE_STRATEGY,
    THEME_COMPARABLE_SALES,
    THEME_GEOGRAPHIC_GROWTH,
    THEME_OPERATING_MARGIN,
    THEME_PRODUCTIVITY,
    THEME_STORE_EXPANSION,
    HistoricalStrategyData,
    HistoricalStrategyDisclosure,
    disclosure_locator,
    load_strategy_disclosures,
)
from modeler.data.standardized_io import standardized_from_payload, standardized_to_payload
from modeler.data.interface import (
    HistoricalManagementKpiDeferredDisagreement,
    HistoricalManagementKpiDeferredMember,
)
from modeler.engine.component_catalog import (
    REVENUE_DRIVER_COMPSALES_DIFFERENCE_FAMILY_ID,
    REVENUE_DRIVER_COMPSALES_FAMILY_ID,
    REVENUE_DRIVER_COMPONENT_CATALOG,
    REVENUE_DRIVER_GEO_CONTRIBUTION_FAMILY_ID,
    REVENUE_DRIVER_REVENUE_GROWTH_FAMILY_ID,
    REVENUE_DRIVER_RPS_FAMILY_ID,
    REVENUE_DRIVER_SHEET_NAME,
    REVENUE_DRIVER_STORE_DIFFERENCE_FAMILY_ID,
    REVENUE_DRIVER_STORE_GROWTH_FAMILY_ID,
    SemanticCellRef,
    expand_revenue_driver_specs,
    is_operating_kpi_source_identity,
    resolve_revenue_driver_link_formula,
    revenue_driver_component_id,
)
from modeler.workbook import ReferenceModelBuilder
from modeler.ingestion.management_kpi_identity import (
    FAMILY_COMPARABLE_SALES_GROWTH,
    FAMILY_SALES_PER_SQUARE_FOOT,
    POP_COMPANY_OPERATED_STORES,
    POP_STORES_AND_DTC,
    POP_STORES_AND_ECOMMERCE,
)
from modeler.line_resolver import MissingLineError
from modeler.period_axis import PeriodAxisError
from composer.overview import (
    DEFERRED_SPSF_LINK,
    PROFESSIONAL_FALLBACK,
    UNTESTED_INITIATIVES,
    WHAT_HISTORY_ESTABLISHES,
)
from composer.revenue_driver import (
    ADDITIONAL_SPSF,
    FAILED_COMPSALES,
    FAILED_SPSF_GROWTH,
    FAILED_STORE_GROWTH,
    SCOPE_NOTE,
)
from director.driver_assessment import (
    complete_historical_strategy_synthesis as compute_historical_strategy_synthesis,
    complete_revenue_driver_analysis as compute_revenue_driver_analysis,
)
from interpreter.revenue_driver import (
    VERDICT_CONTRADICTED,
    VERDICT_INSUFFICIENT,
    VERDICT_MIXED,
    VERDICT_SUPPORTED,
)
from modeler.revenue_driver import revenue_driver_applicable, strategy_synthesis_applicable
from modeler.tests.test_capex import P1, P2, _tiny
from modeler.tests.test_historical_segment import _corp_values, _snapshot
from composer.tests.test_learner_ready_presentation import (
    _assert_answer_key_no_yellow,
    assert_bav_has_no_exercise_framing,
)
from modeler.tests.test_geographic_segment_workbook import _geo_tiny
from modeler.tests.test_operating_kpi_facts import _kpi_model_observation
from modeler.tests.test_operating_kpi_management_history import _compsales, _spsf
from modeler.tests.test_operating_kpi_relationships import _fin_with_relationship
from legacy.trainer.checker import check_workbook
from modeler.tests.test_normalization import _inject_formula_and_cached_value
from modeler.semantic_io import load_semantic_map, parse_cell_ref
from legacy.trainer.derive import build_training_workbook, derive_trainer_workbook

ROOT = Path(__file__).resolve().parents[2]
FR_JSON = ROOT / "build" / "input" / "fast_retailing" / "reconciled" / "standardized.json"
LULU_JSON = ROOT / "core" / "tests" / "fixtures" / "ordinary_reconcile" / "lululemon" / "standardized.json"
LULU_DISCLOSURES = (
    ROOT / "core" / "tests" / "fixtures" / "strategy" / "lululemon_management_disclosures.json"
)
P0 = date(2023, 12, 31)


def _disclosure(
    theme: str,
    *,
    role: str = ROLE_STRATEGY,
    text: str = "We open stores.",
    period: date = P2,
) -> HistoricalStrategyDisclosure:
    return HistoricalStrategyDisclosure(
        theme=theme,
        role=role,
        text=text,
        period=period,
        source_file="example.pdf",
        page_reference="Form 10-K p. 1",
        section="Item 1",
    )


def _with_strategy(fin, *disclosures: HistoricalStrategyDisclosure):
    fin.historical_strategy = HistoricalStrategyData(disclosures=disclosures)
    return fin


def _store_fin(
    counts: dict[date, float],
    revenue: dict[date, float | None],
    *disclosures: HistoricalStrategyDisclosure,
    management=None,
):
    observations = tuple(
        _kpi_model_observation(period, value) for period, value in counts.items()
    )
    fin = _fin_with_relationship(
        *observations,
        revenue=revenue,
        extra_periods=list(counts),
        management=management,
    )
    return _with_strategy(fin, *disclosures)


def _deferred_member(
    *,
    locator: str,
    definition_text: str,
    presentation_role: str = "current",
    extraction_document: str = "example.json",
    page_reference: str = "Form 10-K p. 10",
    physical_page_mapping: str = "10",
    reported_value: float | None = 1580.0,
) -> HistoricalManagementKpiDeferredMember:
    return HistoricalManagementKpiDeferredMember(
        locator=locator,
        extraction_document=extraction_document,
        page_reference=page_reference,
        physical_page_mapping=physical_page_mapping,
        presentation_role=presentation_role,
        definition_text=definition_text,
        population=POP_COMPANY_OPERATED_STORES,
        unit="USD_per_square_foot",
        basis="reported",
        calendar_week_adjustment="included",
        calendar_reporting_basis="52_week",
        reported_value=reported_value,
    )


def _deferred_disagreement(
    *,
    family: str = FAMILY_SALES_PER_SQUARE_FOOT,
    period: date = P2,
    reasons: tuple[str, ...] = ("definition_mismatch", "ordinary_disagreement"),
    members: tuple[HistoricalManagementKpiDeferredMember, ...] | None = None,
) -> HistoricalManagementKpiDeferredDisagreement:
    return HistoricalManagementKpiDeferredDisagreement(
        family=family,
        period=period,
        reasons=list(reasons),
        members=list(
            members
            or (
                _deferred_member(
                    locator=f"{family}:current:{period.isoformat()}",
                    definition_text="average during the year",
                    presentation_role="current",
                ),
                _deferred_member(
                    locator=f"{family}:prior:{period.isoformat()}",
                    definition_text="average ending square footage",
                    presentation_role="prior",
                    page_reference="Form 10-K p. 11",
                    physical_page_mapping="11",
                ),
            )
        ),
    )


def _attach_deferred(fin, *items: HistoricalManagementKpiDeferredDisagreement):
    data = fin.historical_operating_kpis
    assert data is not None
    data.deferred_disagreements = list(items)
    return fin


def _synthesis_record(synthesis):
    return {
        "lead": synthesis.lead,
        "productivity_gap": synthesis.productivity_gap,
        "untested": synthesis.untested,
        "limits": synthesis.limits,
        "navigation": synthesis.navigation,
        "interpretations": tuple(
            (
                item.theme,
                item.heading,
                item.management_statement,
                item.finding,
                item.inference,
                item.supporting_schedules,
                item.counterexample,
            )
            for item in synthesis.interpretations
        ),
    }


def test_director_sequences_historical_strategy_interpretation_before_composition(
    monkeypatch,
):
    from composer.overview import (
        compute_historical_strategy_synthesis as compose,
    )
    from director import driver_assessment as director_mod
    from interpreter.historical_strategy import interpret_historical_strategy

    fin = _store_fin(
        {P0: 10, P1: 12, P2: 15},
        {P0: 100.0, P1: 130.0, P2: 160.0},
        _disclosure(THEME_STORE_EXPANSION),
    )
    analysis = director_mod.complete_revenue_driver_analysis(fin)
    order: list[str] = []
    captured: dict[str, object] = {}

    def tracking_interpret(tests):
        order.append("interpret")
        judgment = interpret_historical_strategy(tests)
        captured["judgment"] = judgment
        return judgment

    def tracking_word(financials, analysis_arg, judgment=None):
        order.append("compose")
        captured["passed_judgment"] = judgment
        captured["passed_analysis"] = analysis_arg
        return compose(financials, analysis_arg, judgment)

    def boom_recompute(*args, **kwargs):
        raise AssertionError("must not recompute supplied analysis")

    monkeypatch.setattr(director_mod, "interpret_historical_strategy", tracking_interpret)
    monkeypatch.setattr(director_mod, "word_historical_strategy", tracking_word)
    monkeypatch.setattr(
        director_mod, "complete_revenue_driver_analysis", boom_recompute
    )

    synthesis = director_mod.complete_historical_strategy_synthesis(fin, analysis)
    assert order == ["interpret", "compose"]
    assert captured["passed_judgment"] is captured["judgment"]
    assert captured["passed_analysis"] is analysis
    assert "historical growth pattern" in synthesis.lead


def test_director_obtains_analysis_when_omitted_then_interprets(monkeypatch):
    from composer.overview import (
        compute_historical_strategy_synthesis as compose,
    )
    from director import driver_assessment as director_mod
    from interpreter.historical_strategy import interpret_historical_strategy

    fin = _store_fin(
        {P0: 10, P1: 12, P2: 15},
        {P0: 100.0, P1: 130.0, P2: 160.0},
        _disclosure(THEME_STORE_EXPANSION),
    )
    real_complete = director_mod.complete_revenue_driver_analysis
    calls: list[object] = []
    order: list[str] = []

    def tracking_complete(financials, periods=None):
        calls.append(financials)
        if periods is None:
            return real_complete(financials)
        return real_complete(financials, periods)

    def tracking_interpret(tests):
        order.append("interpret")
        return interpret_historical_strategy(tests)

    def tracking_word(financials, analysis_arg, judgment=None):
        order.append("compose")
        assert analysis_arg is not None
        assert judgment is not None
        return compose(financials, analysis_arg, judgment)

    monkeypatch.setattr(
        director_mod, "complete_revenue_driver_analysis", tracking_complete
    )
    monkeypatch.setattr(director_mod, "interpret_historical_strategy", tracking_interpret)
    monkeypatch.setattr(director_mod, "word_historical_strategy", tracking_word)

    synthesis = director_mod.complete_historical_strategy_synthesis(fin)
    assert len(calls) == 1
    assert order == ["interpret", "compose"]
    assert "historical growth pattern" in synthesis.lead


def test_historical_strategy_synthesis_records_match_across_orchestration_paths():
    from director.driver_assessment import (
        complete_historical_strategy_synthesis,
        complete_revenue_driver_analysis,
    )
    from interpreter.historical_strategy import interpret_historical_strategy
    from composer.overview import compute_historical_strategy_synthesis as compose

    supported = _store_fin(
        {P0: 10, P1: 12, P2: 15},
        {P0: 100.0, P1: 130.0, P2: 160.0},
        _disclosure(THEME_STORE_EXPANSION, role=ROLE_OBJECTIVE, text="We plan to open stores."),
    )
    mixed = _store_fin(
        {P0: 10, P1: 12, P2: 15},
        {P0: 100.0, P1: 130.0, P2: 120.0},
        _disclosure(THEME_STORE_EXPANSION),
    )
    contradicted = _store_fin(
        {P1: 10, P2: 12},
        {P1: 130.0, P2: 100.0},
        _disclosure(THEME_STORE_EXPANSION),
    )
    insufficient = _with_strategy(
        _tiny(with_payments=False),
        _disclosure(THEME_STORE_EXPANSION),
    )
    geo = _geo_tiny(
        _snapshot(P1, values=_corp_values(rev=(80.0, 25.0, 15.0))),
        _snapshot(P2, values=_corp_values(rev=(110.0, 15.0, 15.0))),
    )
    _with_strategy(geo, _disclosure(THEME_GEOGRAPHIC_GROWTH))
    deferred = _attach_deferred(
        _store_fin(
            {P0: 10, P1: 12, P2: 15},
            {P0: 100.0, P1: 130.0, P2: 160.0},
            _disclosure(THEME_PRODUCTIVITY, role=ROLE_OPERATING_USE),
            management=[_spsf(period=P2, value=1500)],
        ),
        _deferred_disagreement(period=P1),
    )

    for fin in (supported, mixed, contradicted, insufficient, geo, deferred):
        analysis = complete_revenue_driver_analysis(fin)
        judgment = interpret_historical_strategy(analysis.tests)
        via_composer = compose(fin, analysis, judgment)
        via_supplied = complete_historical_strategy_synthesis(fin, analysis)
        via_omitted = complete_historical_strategy_synthesis(fin)
        via_facade = compute_historical_strategy_synthesis(fin, analysis)
        assert _synthesis_record(via_supplied) == _synthesis_record(via_composer)
        assert _synthesis_record(via_omitted) == _synthesis_record(via_composer)
        assert _synthesis_record(via_facade) == _synthesis_record(via_composer)

    supported_row = compose(
        supported,
        complete_revenue_driver_analysis(supported),
        interpret_historical_strategy(complete_revenue_driver_analysis(supported).tests),
    ).interpretations[0]
    assert "descriptively consistent" in supported_row.inference
    mixed_lead = complete_historical_strategy_synthesis(mixed).lead
    assert "mixed across periods or segments" in mixed_lead
    contradicted_lead = complete_historical_strategy_synthesis(contradicted).lead
    assert "contradicts" in contradicted_lead
    insufficient_lead = complete_historical_strategy_synthesis(insufficient).lead
    assert "cannot be treated as a demonstrated historical revenue driver" in insufficient_lead
    geo_synthesis = complete_historical_strategy_synthesis(geo)
    assert "contributed negatively" in geo_synthesis.interpretations[0].inference
    assert "counterexamples" in geo_synthesis.lead
    deferred_synthesis = complete_historical_strategy_synthesis(deferred)
    assert DEFERRED_SPSF_LINK in deferred_synthesis.productivity_gap


DRIVER_FAMILY_IDS = (
    REVENUE_DRIVER_STORE_GROWTH_FAMILY_ID,
    REVENUE_DRIVER_REVENUE_GROWTH_FAMILY_ID,
    REVENUE_DRIVER_STORE_DIFFERENCE_FAMILY_ID,
    REVENUE_DRIVER_COMPSALES_FAMILY_ID,
    REVENUE_DRIVER_COMPSALES_DIFFERENCE_FAMILY_ID,
    REVENUE_DRIVER_RPS_FAMILY_ID,
    REVENUE_DRIVER_GEO_CONTRIBUTION_FAMILY_ID,
)


def _copy_workbooks(trainer: Path, answer: Path, dest: Path) -> tuple[Path, Path]:
    dest.mkdir()
    copied_trainer = dest / trainer.name
    copied_answer = dest / answer.name
    shutil.copy2(trainer, copied_trainer)
    shutil.copy2(answer, copied_answer)
    for sidecar in (
        answer.with_suffix(".component_map.json"),
        answer.with_suffix(".assumptions.json"),
        answer.with_suffix(".trainer.json"),
    ):
        if sidecar.is_file():
            shutil.copy2(sidecar, dest / sidecar.name)
    return copied_trainer, copied_answer


def _assert_readable_driver_layout(sheet) -> None:
    long_cells = 0
    for row in sheet.iter_rows(min_col=1, max_col=2, max_row=sheet.max_row or 1):
        for cell in row:
            value = cell.value
            if not isinstance(value, str) or value.startswith("="):
                continue
            if len(value) < 80:
                continue
            long_cells += 1
            assert cell.alignment.wrap_text is True, cell.coordinate
            height = sheet.row_dimensions[cell.row].height or 15
            assert height > 15, (cell.coordinate, len(value), height)
    assert long_cells >= 1


def test_revenue_then_margin_assessment_order_is_first_name_wins():
    from director.current_build import prepare_company_input, resolve_company
    from core.research.drivers import assemble_drivers_view, _unique_assessments

    fin = prepare_company_input(resolve_company("Lululemon"), Path("/tmp/unused"))
    analysis = compute_revenue_driver_analysis(fin)
    names = [item.name for item in analysis.assessments]
    theme_order = [
        "footprint and intensity identity",
        "comparable-sales coincidence",
        "sales-per-square-foot productivity",
        "geographic revenue reconstruction",
    ]
    positions = [names.index(item) for item in theme_order]
    assert positions == sorted(positions)
    margin_names = [
        "component operating-margin identity",
        "component operating-margin contributions",
        "gross-profit amount bridge",
        "impairment or asset-related charges",
        "mix, markdowns, freight, costs, or leverage",
        "latest adjacent operating-margin movement",
    ]
    margin_positions = [names.index(item) for item in margin_names]
    assert margin_positions == sorted(margin_positions)
    assert names.index(theme_order[0]) < names.index(margin_names[0])
    view = assemble_drivers_view(fin, "Lululemon")
    unique_names = [item.name for item in view.assessments]
    assert unique_names == list(dict.fromkeys(unique_names))
    duplicated = _unique_assessments(analysis.assessments + analysis.assessments)
    assert [item.name for item in duplicated] == unique_names or [
        item.name for item in duplicated
    ] == list(dict.fromkeys(names))
