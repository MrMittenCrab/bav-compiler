"""Step 9M.2 — Lululemon real-company benchmark baseline acceptance."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import subprocess
import sys
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from modeler.data.standardized_io import (
    standardized_from_payload,
    standardized_to_payload,
)
from extractor.data.filing_json import load_extracted_filing
from director.ingestion.filing_validator import validate_extracted_filing
from modeler.classification import (
    UnclassifiedBalanceSheetLineError,
    classify_balance_sheet_line,
    is_balance_sheet_subtotal,
)
from modeler.workbook import (
    EARNINGS_QUALITY_SHEET,
    SOURCE_START_ROW,
    ReferenceModelBuilder,
)
from modeler.capex import (
    capex_applicable,
    capex_availability,
    cash_after_ppe_capex_applicable,
    compute_capex_series,
    resolve_capex_source,
    resolve_operating_cash_source,
)
from modeler.deferred_tax import (
    compute_deferred_tax_series,
    deferred_tax_applicable,
    deferred_tax_availability,
    resolve_deferred_tax_sources,
)
from modeler.financial_math import compute_anchor
from modeler.fixed_asset import fixed_asset_applicable, fixed_asset_availability
from modeler.lease_rou import (
    compute_lease_rou_series,
    lease_rou_applicable,
    lease_rou_availability,
    resolve_lease_rou_source,
)
from modeler.line_resolver import resolve_line, workbook_row_for
from modeler.period_axis import canonical_fiscal_periods
from modeler.ratio_values import SOURCE_UNAVAILABLE
from modeler.source_values import required_period_value
from legacy.trainer.checker import check_workbook
from modeler.semantic_io import load_semantic_map, parse_cell_ref
from legacy.trainer.derive import build_training_workbook

ROOT = Path(__file__).resolve().parents[2]
BENCH = ROOT / "build" / "input" / "lululemon"
EXTRACTED = BENCH / "extracted"
SOURCE = BENCH / "source"
RECONCILED = ROOT / "core" / "tests" / "fixtures" / "ordinary_reconcile" / "lululemon"
STD_JSON = RECONCILED / "standardized.json"
PROV_JSON = RECONCILED / "provenance.json"
CONFLICTS_JSON = RECONCILED / "conflicts.json"

SOURCE_PDFS = {
    2022: (
        "LULU_FY2022_Annual_Report.pdf",
        "b344d1e7a710259fa06f88773dee0b3827334820ce2b881fe6b95ca2ae275e4e",
        4913067,
    ),
    2023: (
        "LULU_FY2023_Annual_Report.pdf",
        "cd47ea251d608d06a3e58b5d782f2d41d5a231a994d2f7993267a430cb13c0f1",
        5848446,
    ),
    2024: (
        "LULU_FY2024_Annual_Report.pdf",
        "9268fd530db162babdd1ec4363cf388ebce57125d83b7e097aba6f98ba0ca7ec",
        5953217,
    ),
    2025: (
        "LULU_FY2025_Annual_Report.pdf",
        "82e00f900cc912a7d79596409594156b7779c3a193783ea8fecf87bc013c71cc",
        6590658,
    ),
}

EXPECTED_PERIODS = [
    date(2022, 1, 30),
    date(2023, 1, 29),
    date(2024, 1, 28),
    date(2025, 2, 2),
    date(2026, 2, 1),
]
FOUR_PERIOD_AXIS = EXPECTED_PERIODS[1:]

REVENUE_ANCHORS = {
    date(2022, 1, 30): 6256617.0,
    date(2023, 1, 29): 8110518.0,
    date(2024, 1, 28): 9619278.0,
    date(2025, 2, 2): 10588126.0,
    date(2026, 2, 1): 11102600.0,
}

DILUTED_WAS_ANCHORS = {
    date(2022, 1, 30): 130295.0,
    date(2023, 1, 29): 128017.0,
    date(2024, 1, 28): 127060.0,
    date(2025, 2, 2): 123935.0,
    date(2026, 2, 1): 119068.0,
}

BUILD_BLOCKER_LABELS: set[str] = set()


def _load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_four_filings_validate_and_remain_source_bound():
    files = sorted(EXTRACTED.glob("LULU_FY[0-9][0-9][0-9][0-9].json"))
    assert {int(p.stem.replace("LULU_FY", "")) for p in files} == {2022, 2023, 2024, 2025}

    for fiscal_year, (pdf_name, digest, nbytes) in SOURCE_PDFS.items():
        pdf = SOURCE / pdf_name
        assert pdf.is_file()
        data = pdf.read_bytes()
        assert len(data) == nbytes
        assert hashlib.sha256(data).hexdigest() == digest

        extracted = EXTRACTED / f"LULU_FY{fiscal_year}.json"
        filing = load_extracted_filing(extracted)
        assert filing.schema_version == "1.0"
        assert filing.filing.fiscal_year == fiscal_year
        assert filing.filing.source_file == pdf_name
        assert filing.income_statement
        assert filing.balance_sheet
        assert filing.cash_flow

        report = validate_extracted_filing(filing, source_root=SOURCE)
        assert report.ok
        assert report.computed_source_sha256 == digest
        assert report.warnings == ()


def _reconcile_cmd(out: Path) -> list[str]:
    return [
        sys.executable,
        "-m",
        "core",
        "reconcile",
        str(EXTRACTED),
        "--source-root",
        str(SOURCE),
        "-o",
        str(out),
        "--admit-period",
        "2022-01-30",
    ]


def _default_reconcile_cmd(out: Path) -> list[str]:
    return [
        sys.executable,
        "-m",
        "core",
        "reconcile",
        str(EXTRACTED),
        "--source-root",
        str(SOURCE),
        "-o",
        str(out),
    ]


def _subprocess_env() -> dict[str, str]:
    env = dict(os.environ)
    env["PYTHONPATH"] = str(ROOT)
    return env


ARTIFACT_NAMES = ("standardized.json", "provenance.json", "conflicts.json")


def _read_committed_artifacts() -> dict[str, bytes]:
    return {name: (RECONCILED / name).read_bytes() for name in ARTIFACT_NAMES}


# Protected ordinary_reconcile fixtures predate sparse IS component retention.
# Compare the inherited contract without those newly retained face-of-statement
# series; focused tests assert the expanded components separately.
_EXPANDED_IS_COMPONENT_CONCEPTS = frozenset(
    {
        "impairment_and_restructuring",
        "acquisition_related_expenses",
        "gain_on_disposal_of_assets",
    }
)


def _is_expanded_is_component(entry: dict) -> bool:
    return (
        entry.get("statement") == "income_statement"
        and entry.get("suggested_concept") in _EXPANDED_IS_COMPONENT_CONCEPTS
    )


def _statement_provenance(payload: dict) -> dict:
    """Canonical-comparable provenance: statement selections only."""
    comparable = dict(payload)
    comparable["note_facts"] = []
    comparable.pop("selected_geographic_segment_facts", None)
    comparable.pop("selected_operating_kpi_facts", None)
    comparable["omitted_incomplete_axis"] = [
        item
        for item in comparable.get("omitted_incomplete_axis", [])
        if not _is_expanded_is_component(item)
    ]
    comparable["retained_sparse_axis"] = [
        item
        for item in comparable.get("retained_sparse_axis", [])
        if not _is_expanded_is_component(item)
    ]
    comparable["values"] = {
        key: value
        for key, value in comparable.get("values", {}).items()
        if not _is_expanded_is_component(value)
    }
    return comparable


def _comparable_standardized(payload: dict) -> dict:
    """Canonical-comparable standardized payload: ignore optional handoffs."""
    from datetime import date
    from modeler.data.issuer_fiscal import (
        issuer_fiscal_label,
        issuer_fiscal_years_from_extracted,
    )
    comparable = dict(payload)
    comparable.pop("historical_segment", None)
    comparable.pop("historical_operating_kpis", None)
    comparable["income_statement"] = [
        item
        for item in comparable.get("income_statement", [])
        if item.get("concept") not in _EXPANDED_IS_COMPONENT_CONCEPTS
    ]
    mapping = issuer_fiscal_years_from_extracted(EXTRACTED)
    periods = []
    for item in comparable.get("periods", []):
        period = dict(item)
        end = date.fromisoformat(str(period["end_date"]))
        period["label"] = issuer_fiscal_label(mapping[end])
        periods.append(period)
    comparable["periods"] = periods
    return comparable


def _assert_artifact_sets_match(
    generated: Path,
    expected_bytes: dict[str, bytes],
) -> None:
    """Compare generated artifacts to expected bytes; never write expected paths."""
    mismatches: list[str] = []
    for name in ARTIFACT_NAMES:
        actual = (generated / name).read_bytes()
        if name == "provenance.json":
            try:
                actual_payload = json.loads(actual)
                expected_payload = json.loads(expected_bytes[name])
            except json.JSONDecodeError:
                mismatches.append(name)
                continue
            if _statement_provenance(actual_payload) != _statement_provenance(
                expected_payload
            ):
                mismatches.append(name)
            continue
        if name == "standardized.json":
            try:
                actual_payload = json.loads(actual)
                expected_payload = json.loads(expected_bytes[name])
            except json.JSONDecodeError:
                mismatches.append(name)
                continue
            if _comparable_standardized(actual_payload) != _comparable_standardized(
                expected_payload
            ):
                mismatches.append(name)
            continue
        if actual != expected_bytes[name]:
            mismatches.append(name)
    if mismatches:
        raise AssertionError(
            "reconciliation artifact mismatch: " + ", ".join(mismatches)
        )


def _run_reconcile_pass(out: Path) -> None:
    completed = subprocess.run(
        _reconcile_cmd(out),
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
        env=_subprocess_env(),
    )
    assert "overlap_conflicts=" in completed.stdout


def _assert_inter_run_artifacts_match(out1: Path, out2: Path) -> None:
    for name in ARTIFACT_NAMES:
        assert (out1 / name).read_bytes() == (out2 / name).read_bytes()


def _assert_committed_artifacts_unchanged(committed_before: dict[str, bytes]) -> None:
    committed_after = _read_committed_artifacts()
    mutated = [
        name
        for name in ARTIFACT_NAMES
        if committed_after[name] != committed_before[name]
    ]
    if mutated:
        raise AssertionError(
            "committed reconciliation artifacts mutated: " + ", ".join(mutated)
        )


def _guarded_deterministic_reconcile(
    tmp_path: Path,
    *,
    committed_before: dict[str, bytes] | None = None,
    run_pass=_run_reconcile_pass,
    assert_inter_run=_assert_inter_run_artifacts_match,
    assert_baseline=_assert_artifact_sets_match,
    read_committed=_read_committed_artifacts,
) -> None:
    """Run both reconcile passes under try; always re-verify committed bytes in finally."""
    if committed_before is None:
        committed_before = read_committed()
    out1 = tmp_path / "r1"
    out2 = tmp_path / "r2"
    error: BaseException | None = None
    try:
        run_pass(out1)
        run_pass(out2)
        assert_inter_run(out1, out2)
        assert_baseline(out1, committed_before)
        assert_baseline(out2, committed_before)
    except BaseException as exc:
        error = exc
    finally:
        try:
            after = read_committed()
            mutated = [
                name
                for name in ARTIFACT_NAMES
                if after[name] != committed_before[name]
            ]
            if mutated:
                raise AssertionError(
                    "committed reconciliation artifacts mutated: "
                    + ", ".join(mutated)
                )
        except AssertionError:
            raise
        else:
            if error is not None:
                raise error


COMMON_STOCK_VALUES = {
    date(2022, 1, 30): 616.0,
    date(2023, 1, 29): 611.0,
    date(2024, 1, 28): 606.0,
    date(2025, 2, 2): 581.0,
    date(2026, 2, 1): 557.0,
}


CAPEX_REPORTED_PAYMENTS = {
    date(2022, 1, 30): -394502.0,
    date(2023, 1, 29): -638657.0,
    date(2024, 1, 28): -651865.0,
    date(2025, 2, 2): -689232.0,
    date(2026, 2, 1): -680802.0,
}


SBC_REPORTED = {
    date(2022, 1, 30): 69137.0,
    date(2023, 1, 29): 78075.0,
    date(2024, 1, 28): 93560.0,
    date(2025, 2, 2): 90011.0,
    date(2026, 2, 1): 62203.0,
}
CFO_LESS_SBC = {
    date(2022, 1, 30): 1319971.0,
    date(2023, 1, 29): 888388.0,
    date(2024, 1, 28): 2202604.0,
    date(2025, 2, 2): 2182702.0,
    date(2026, 2, 1): 1540274.0,
}


ACQUISITION_REPORTED = {
    date(2022, 1, 30): 0.0,
    date(2023, 1, 29): 0.0,
    date(2024, 1, 28): 0.0,
    date(2025, 2, 2): -154146.0,
    date(2026, 2, 1): 0.0,
}
ACQUISITION_OUTFLOW = {
    date(2022, 1, 30): 0.0,
    date(2023, 1, 29): 0.0,
    date(2024, 1, 28): 0.0,
    date(2025, 2, 2): 154146.0,
    date(2026, 2, 1): 0.0,
}
CASH_AFTER_PPE_CAPEX_AND_ACQUISITIONS = {
    date(2022, 1, 30): 994606.0,
    date(2023, 1, 29): 327806.0,
    date(2024, 1, 28): 1644299.0,
    date(2025, 2, 2): 1429335.0,
    date(2026, 2, 1): 921675.0,
}


REPURCHASE_REPORTED = {
    date(2022, 1, 30): -812602.0,
    date(2023, 1, 29): -444001.0,
    date(2024, 1, 28): -558652.0,
    date(2025, 2, 2): -1636879.0,
    date(2026, 2, 1): -1178349.0,
}
REPURCHASE_OUTFLOW = {
    date(2022, 1, 30): 812602.0,
    date(2023, 1, 29): 444001.0,
    date(2024, 1, 28): 558652.0,
    date(2025, 2, 2): 1636879.0,
    date(2026, 2, 1): 1178349.0,
}
CASH_AFTER_PPE_CAPEX_ACQUISITIONS_AND_REPURCHASES = {
    date(2022, 1, 30): 182004.0,
    date(2023, 1, 29): -116195.0,
    date(2024, 1, 28): 1085647.0,
    date(2025, 2, 2): -207544.0,
    date(2026, 2, 1): -256674.0,
}


CASH_MOVEMENT_FROM_FLOWS = {
    date(2022, 1, 30): 109354.0,
    date(2023, 1, 29): -105004.0,
    date(2024, 1, 28): 1089104.0,
    date(2025, 2, 2): -259635.0,
    date(2026, 2, 1): -177134.0,
}


ROU_BALANCES = {
    date(2022, 1, 30): 803543.0,
    date(2023, 1, 29): 969419.0,
    date(2024, 1, 28): 1265610.0,
    date(2025, 2, 2): 1416256.0,
    date(2026, 2, 1): 1630181.0,
}
DTA_BALANCES = {
    date(2022, 1, 30): 6091.0,
    date(2023, 1, 29): 6402.0,
    date(2024, 1, 28): 9176.0,
    date(2025, 2, 2): 17085.0,
    date(2026, 2, 1): 24037.0,
}
DTL_BALANCES = {
    date(2022, 1, 30): 53352.0,
    date(2023, 1, 29): 55084.0,
    date(2024, 1, 28): 29522.0,
    date(2025, 2, 2): 98188.0,
    date(2026, 2, 1): 52278.0,
}
NET_DT_POSITIONS = {
    date(2022, 1, 30): -47261.0,
    date(2023, 1, 29): -48682.0,
    date(2024, 1, 28): -20346.0,
    date(2025, 2, 2): -81103.0,
    date(2026, 2, 1): -28241.0,
}
PRIOR_LULULEMON_SPECS = 453
LEASE_DT_LULULEMON_SPECS = 486


def _fill_rgb(cell) -> str:
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return ""
    color = fill.fgColor.rgb or fill.start_color.rgb or ""
    return str(color).upper().lstrip("0")[-6:] if color else ""


def _count_source_unavailable(path: Path) -> int:
    wb = load_workbook(path, data_only=False)
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == SOURCE_UNAVAILABLE:
                    n += 1
    wb.close()
    return n


def test_no_lulu_specific_production_branch():
    """Production engine must stay generic — no ticker/issuer hard-codes."""
    core_root = ROOT / "core"
    offenders: list[str] = []
    for path in core_root.rglob("*.py"):
        if "tests" in path.parts:
            continue
        text = path.read_text(encoding="utf-8")
        for token in ("LULU", "lululemon", "Lululemon"):
            if token in text:
                offenders.append(f"{path.relative_to(ROOT)}:{token}")
    assert offenders == []


def test_committed_reconciled_hashes_are_stable():
    """Lock measured baseline artifact digests for Step 9M.2.4.1."""
    assert _sha256(STD_JSON) == (
        "a3568c29e883c8ba57af23da7b4286641a3c5f929af311e2e9593c5f63ea2287"
    )
    assert _sha256(CONFLICTS_JSON) == (
        "d8a33012f6ea73126ac4e2ece3613e7011c11cb2b581745d8c3563e3c2e978e0"
    )
    # provenance is large; lock size + digest together
    assert PROV_JSON.stat().st_size == 699438
    assert _sha256(PROV_JSON) == (
        "6799371215e02c888b3a5f687637b38dba4840bd1548cb1860a253f7a699cb12"
    )
