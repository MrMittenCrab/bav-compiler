"""Disclosure-led historical revenue-driver tests and workbook presentation."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path
import shutil

import pytest
from openpyxl import load_workbook

from bav.extractor.data.historical_strategy import (
    ROLE_OBJECTIVE,
    ROLE_OPERATING_USE,
    ROLE_STRATEGY,
    THEME_COMPARABLE_SALES,
    THEME_GEOGRAPHIC_GROWTH,
    THEME_OPERATING_MARGIN,
    THEME_PRODUCTIVITY,
    THEME_STORE_EXPANSION,
    HistoricalStrategyData,
    HistoricalStrategyDisclosure,
    disclosure_locator,
    load_strategy_disclosures,
)
from bav.modeler.data.standardized_io import standardized_from_payload, standardized_to_payload
from bav.modeler.data.interface import (
    HistoricalManagementKpiDeferredDisagreement,
    HistoricalManagementKpiDeferredMember,
)
from bav.modeler.engine.component_catalog import (
    REVENUE_DRIVER_COMPSALES_DIFFERENCE_FAMILY_ID,
    REVENUE_DRIVER_COMPSALES_FAMILY_ID,
    REVENUE_DRIVER_COMPONENT_CATALOG,
    REVENUE_DRIVER_GEO_CONTRIBUTION_FAMILY_ID,
    REVENUE_DRIVER_REVENUE_GROWTH_FAMILY_ID,
    REVENUE_DRIVER_RPS_FAMILY_ID,
    REVENUE_DRIVER_SHEET_NAME,
    REVENUE_DRIVER_STORE_DIFFERENCE_FAMILY_ID,
    REVENUE_DRIVER_STORE_GROWTH_FAMILY_ID,
    SemanticCellRef,
    expand_revenue_driver_specs,
    is_operating_kpi_source_identity,
    resolve_revenue_driver_link_formula,
    revenue_driver_component_id,
)
from bav.modeler.workbook import ReferenceModelBuilder
from bav.modeler.ingestion.management_kpi_identity import (
    FAMILY_COMPARABLE_SALES_GROWTH,
    FAMILY_SALES_PER_SQUARE_FOOT,
    POP_COMPANY_OPERATED_STORES,
    POP_STORES_AND_DTC,
    POP_STORES_AND_ECOMMERCE,
)
from bav.modeler.line_resolver import MissingLineError
from bav.modeler.period_axis import PeriodAxisError
from bav.composer.overview import (
    DEFERRED_SPSF_LINK,
    PROFESSIONAL_FALLBACK,
    UNTESTED_INITIATIVES,
    WHAT_HISTORY_ESTABLISHES,
)
from bav.composer.revenue_driver import (
    ADDITIONAL_SPSF,
    FAILED_COMPSALES,
    FAILED_SPSF_GROWTH,
    FAILED_STORE_GROWTH,
    SCOPE_NOTE,
)
from bav.director.driver_assessment import (
    complete_historical_strategy_synthesis as compute_historical_strategy_synthesis,
    complete_revenue_driver_analysis as compute_revenue_driver_analysis,
)
from bav.inferer.revenue_driver import VERDICT_CONTRADICTED, VERDICT_INSUFFICIENT, VERDICT_MIXED, VERDICT_SUPPORTED
from bav.modeler.revenue_driver import revenue_driver_applicable, strategy_synthesis_applicable
from bav.modeler.tests.test_capex import P1, P2, _tiny
from bav.modeler.tests.test_historical_segment import _corp_values, _snapshot
from bav.composer.tests.test_learner_ready_presentation import (
    _assert_answer_key_no_yellow,
    assert_bav_has_no_exercise_framing,
)
from bav.modeler.tests.test_geographic_segment_workbook import _geo_tiny
from bav.modeler.tests.test_operating_kpi_facts import _kpi_model_observation
from bav.modeler.tests.test_operating_kpi_management_history import _compsales, _spsf
from bav.modeler.tests.test_operating_kpi_relationships import _fin_with_relationship
from legacy.trainer.checker import check_workbook
from bav.modeler.tests.test_normalization import _inject_formula_and_cached_value
from bav.modeler.semantic_io import load_semantic_map, parse_cell_ref
from legacy.trainer.derive import build_training_workbook, derive_trainer_workbook

ROOT = Path(__file__).resolve().parents[3]
FR_JSON = ROOT / "build" / "input" / "fast_retailing" / "reconciled" / "standardized.json"
LULU_JSON = ROOT / "bav" / "modeler" / "tests" / "fixtures" / "ordinary_reconcile" / "lululemon" / "standardized.json"
LULU_DISCLOSURES = (
    ROOT / "bav" / "extractor" / "tests" / "fixtures" / "strategy" / "lululemon_management_disclosures.json"
)
P0 = date(2023, 12, 31)


def _disclosure(
    theme: str,
    *,
    role: str = ROLE_STRATEGY,
    text: str = "We open stores.",
    period: date = P2,
) -> HistoricalStrategyDisclosure:
    return HistoricalStrategyDisclosure(
        theme=theme,
        role=role,
        text=text,
        period=period,
        source_file="example.pdf",
        page_reference="Form 10-K p. 1",
        section="Item 1",
    )


def _with_strategy(fin, *disclosures: HistoricalStrategyDisclosure):
    fin.historical_strategy = HistoricalStrategyData(disclosures=disclosures)
    return fin


def _store_fin(
    counts: dict[date, float],
    revenue: dict[date, float | None],
    *disclosures: HistoricalStrategyDisclosure,
    management=None,
):
    observations = tuple(
        _kpi_model_observation(period, value) for period, value in counts.items()
    )
    fin = _fin_with_relationship(
        *observations,
        revenue=revenue,
        extra_periods=list(counts),
        management=management,
    )
    return _with_strategy(fin, *disclosures)


def _deferred_member(
    *,
    locator: str,
    definition_text: str,
    presentation_role: str = "current",
    extraction_document: str = "example.json",
    page_reference: str = "Form 10-K p. 10",
    physical_page_mapping: str = "10",
    reported_value: float | None = 1580.0,
) -> HistoricalManagementKpiDeferredMember:
    return HistoricalManagementKpiDeferredMember(
        locator=locator,
        extraction_document=extraction_document,
        page_reference=page_reference,
        physical_page_mapping=physical_page_mapping,
        presentation_role=presentation_role,
        definition_text=definition_text,
        population=POP_COMPANY_OPERATED_STORES,
        unit="USD_per_square_foot",
        basis="reported",
        calendar_week_adjustment="included",
        calendar_reporting_basis="52_week",
        reported_value=reported_value,
    )


def _deferred_disagreement(
    *,
    family: str = FAMILY_SALES_PER_SQUARE_FOOT,
    period: date = P2,
    reasons: tuple[str, ...] = ("definition_mismatch", "ordinary_disagreement"),
    members: tuple[HistoricalManagementKpiDeferredMember, ...] | None = None,
) -> HistoricalManagementKpiDeferredDisagreement:
    return HistoricalManagementKpiDeferredDisagreement(
        family=family,
        period=period,
        reasons=list(reasons),
        members=list(
            members
            or (
                _deferred_member(
                    locator=f"{family}:current:{period.isoformat()}",
                    definition_text="average during the year",
                    presentation_role="current",
                ),
                _deferred_member(
                    locator=f"{family}:prior:{period.isoformat()}",
                    definition_text="average ending square footage",
                    presentation_role="prior",
                    page_reference="Form 10-K p. 11",
                    physical_page_mapping="11",
                ),
            )
        ),
    )


def _attach_deferred(fin, *items: HistoricalManagementKpiDeferredDisagreement):
    data = fin.historical_operating_kpis
    assert data is not None
    data.deferred_disagreements = list(items)
    return fin


def test_store_expansion_supported_mixed_contradicted_and_insufficient():
    supported = _store_fin(
        {P0: 10, P1: 12, P2: 15},
        {P0: 100.0, P1: 130.0, P2: 160.0},
        _disclosure(THEME_STORE_EXPANSION, role=ROLE_OBJECTIVE, text="We plan to open stores."),
    )
    result = compute_revenue_driver_analysis(supported)
    assert result.scope_note == SCOPE_NOTE
    assert result.calculation_kind == "analyst-derived"
    test = result.tests[0]
    assert test.theme == THEME_STORE_EXPANSION
    assert test.verdict == VERDICT_SUPPORTED
    assert test.sample_size == 2
    assert test.failed_requirement == ""
    assert any(item.role == ROLE_OBJECTIVE for item in test.disclosures)
    assert "not treated as achieved historical outcomes" in " ".join(test.limitations)
    assert "descriptive" in " ".join(test.limitations).lower()
    assert "divided by company-operated stores" in " ".join(test.identity_notes).lower()
    notes = " ".join(item.note for item in test.observations)
    assert "Both consolidated revenue and company-operated store counts grew." in notes

    mixed = _store_fin(
        {P0: 10, P1: 12, P2: 15},
        {P0: 100.0, P1: 130.0, P2: 120.0},
        _disclosure(THEME_STORE_EXPANSION),
    )
    mixed_test = compute_revenue_driver_analysis(mixed).tests[0]
    assert mixed_test.verdict == VERDICT_MIXED
    assert mixed_test.sample_size == 2
    assert any(
        "did not" in item.note for item in mixed_test.observations if item.consistent is False
    )

    contradicted = _store_fin(
        {P1: 10, P2: 12},
        {P1: 130.0, P2: 100.0},
        _disclosure(THEME_STORE_EXPANSION),
    )
    contradicted_test = compute_revenue_driver_analysis(contradicted).tests[0]
    assert contradicted_test.verdict == VERDICT_CONTRADICTED
    assert contradicted_test.sample_size == 1

    missing_history = _with_strategy(
        _tiny(with_payments=False),
        _disclosure(THEME_STORE_EXPANSION),
    )
    missing_test = compute_revenue_driver_analysis(missing_history).tests[0]
    assert missing_test.verdict == VERDICT_INSUFFICIENT
    assert missing_test.sample_size == 0
    assert missing_test.failed_requirement == FAILED_STORE_GROWTH


def test_productivity_insufficient_without_adjacent_spsf_growth():
    fin = _store_fin(
        {P1: 10, P2: 12},
        {P1: 100.0, P2: 130.0},
        _disclosure(THEME_PRODUCTIVITY, role=ROLE_OPERATING_USE),
    )
    test = compute_revenue_driver_analysis(fin).tests[0]
    assert test.verdict == VERDICT_INSUFFICIENT
    assert test.failed_requirement == FAILED_SPSF_GROWTH
    assert ADDITIONAL_SPSF in test.additional_evidence
    assert "divided by company-operated stores" in " ".join(test.identity_notes).lower()
    joined_limits = " ".join(test.limitations)
    assert "2023-01-29" not in joined_limits
    assert "definition disagreement" not in joined_limits.lower()


def test_comparable_sales_limitations_follow_admitted_populations():
    fin = _store_fin(
        {P1: 10, P2: 12},
        {P1: 100.0, P2: 130.0},
        _disclosure(THEME_COMPARABLE_SALES),
        management=[
            _compsales(
                period=P1,
                value=4.0,
                geography="",
                basis="reported",
                population=POP_COMPANY_OPERATED_STORES,
            ),
            _compsales(
                period=P2,
                value=5.0,
                geography="global",
                basis="reported",
                population=POP_STORES_AND_ECOMMERCE,
            ),
            _compsales(
                period=P1,
                value=6.0,
                geography="global",
                basis="reported",
                population=POP_STORES_AND_DTC,
            ),
        ],
    )
    test = compute_revenue_driver_analysis(fin).tests[0]
    joined = " ".join(test.limitations)
    assert POP_COMPANY_OPERATED_STORES in joined
    assert POP_STORES_AND_ECOMMERCE in joined
    assert POP_STORES_AND_DTC in joined
    assert P1.isoformat() in joined
    assert P2.isoformat() in joined
    assert "FY2022" not in joined
    assert "store-only and later stores-plus-DTC" not in joined


def test_productivity_limitations_follow_missing_spsf_periods():
    fin = _store_fin(
        {P0: 10, P1: 12, P2: 15},
        {P0: 100.0, P1: 130.0, P2: 160.0},
        _disclosure(THEME_PRODUCTIVITY),
        management=[
            _spsf(period=P1, value=1600),
            _spsf(period=P2, value=1500),
        ],
    )
    test = compute_revenue_driver_analysis(fin).tests[0]
    joined = " ".join(test.limitations)
    assert P0.isoformat() in joined
    assert "not bridged" in joined.lower()
    assert "2023-01-29" not in joined
    assert "definition disagreement" not in joined.lower()
    assert "ordinary_disagreement" not in joined


def test_productivity_discloses_evidence_derived_deferred_disagreement():
    period = P1
    disagreement = _deferred_disagreement(period=period)
    fin = _attach_deferred(
        _store_fin(
            {P0: 10, P1: 12, P2: 15},
            {P0: 100.0, P1: 130.0, P2: 160.0},
            _disclosure(THEME_PRODUCTIVITY),
            management=[_spsf(period=P2, value=1500)],
        ),
        disagreement,
    )
    test = compute_revenue_driver_analysis(fin).tests[0]
    joined = " ".join(test.limitations)
    assert period.isoformat() in joined
    assert "definition_mismatch" in joined
    assert "ordinary_disagreement" in joined
    assert "average during the year" in joined
    assert "average ending square footage" in joined
    assert disagreement.members[0].locator in joined
    assert disagreement.members[1].locator in joined
    assert "example.json" in joined
    assert "Form 10-K p. 10" in joined
    assert "audit-only" in joined.lower()
    assert "not admitted" in joined.lower()
    assert "does not create an admitted observation" in joined
    assert "2023-01-29" not in joined
    round_trip = standardized_from_payload(standardized_to_payload(fin))
    restored = round_trip.historical_operating_kpis.deferred_disagreements
    assert len(restored) == 1
    assert restored[0].period == period
    assert restored[0].reasons == list(disagreement.reasons)
    assert [item.locator for item in restored[0].members] == [
        item.locator for item in disagreement.members
    ]
    assert all(
        item.period != period
        for item in round_trip.historical_operating_kpis.management_observations
        if item.family == FAMILY_SALES_PER_SQUARE_FOOT
    )


def test_geographic_mixed_when_a_segment_subtracts():
    from bav.modeler.tests.test_geographic_segment_workbook import _geo_tiny

    fin = _geo_tiny(
        _snapshot(P1, values=_corp_values(rev=(80.0, 25.0, 15.0))),
        _snapshot(P2, values=_corp_values(rev=(110.0, 15.0, 15.0))),
    )
    _with_strategy(fin, _disclosure(THEME_GEOGRAPHIC_GROWTH))
    test = compute_revenue_driver_analysis(fin).tests[0]
    assert test.theme == THEME_GEOGRAPHIC_GROWTH
    assert test.verdict == VERDICT_MIXED
    assert test.sample_size == 1
    assert "not organic, constant-currency, or causal" in " ".join(test.limitations)
    assert any("negatively" in item.note for item in test.observations)
    assert any(item.mix_conflict and item.counterexample for item in test.observations)
    assert "contributed negatively" in test.finding


def test_strategy_synthesis_links_deferred_spsf_without_promoting_it():
    period = P1
    disagreement = _deferred_disagreement(period=period)
    fin = _attach_deferred(
        _store_fin(
            {P0: 10, P1: 12, P2: 15},
            {P0: 100.0, P1: 130.0, P2: 160.0},
            _disclosure(THEME_PRODUCTIVITY, role=ROLE_OPERATING_USE),
            management=[_spsf(period=P2, value=1500)],
        ),
        disagreement,
    )
    synthesis = compute_historical_strategy_synthesis(fin)
    assert DEFERRED_SPSF_LINK in synthesis.productivity_gap
    assert "sample size 0" in synthesis.productivity_gap
    assert disagreement.members[0].locator not in synthesis.productivity_gap
    assert disagreement.members[0].definition_text not in synthesis.productivity_gap
    assert "average during the year" not in synthesis.lead
    productivity = next(
        item for item in synthesis.interpretations if item.theme == THEME_PRODUCTIVITY
    )
    assert disagreement.members[0].locator not in productivity.finding
    assert disagreement.members[0].locator not in productivity.inference


def _synthesis_record(synthesis):
    return {
        "lead": synthesis.lead,
        "productivity_gap": synthesis.productivity_gap,
        "untested": synthesis.untested,
        "limits": synthesis.limits,
        "navigation": synthesis.navigation,
        "interpretations": tuple(
            (
                item.theme,
                item.heading,
                item.management_statement,
                item.finding,
                item.inference,
                item.supporting_schedules,
                item.counterexample,
            )
            for item in synthesis.interpretations
        ),
    }


DRIVER_FAMILY_IDS = (
    REVENUE_DRIVER_STORE_GROWTH_FAMILY_ID,
    REVENUE_DRIVER_REVENUE_GROWTH_FAMILY_ID,
    REVENUE_DRIVER_STORE_DIFFERENCE_FAMILY_ID,
    REVENUE_DRIVER_COMPSALES_FAMILY_ID,
    REVENUE_DRIVER_COMPSALES_DIFFERENCE_FAMILY_ID,
    REVENUE_DRIVER_RPS_FAMILY_ID,
    REVENUE_DRIVER_GEO_CONTRIBUTION_FAMILY_ID,
)


def _copy_workbooks(trainer: Path, answer: Path, dest: Path) -> tuple[Path, Path]:
    dest.mkdir()
    copied_trainer = dest / trainer.name
    copied_answer = dest / answer.name
    shutil.copy2(trainer, copied_trainer)
    shutil.copy2(answer, copied_answer)
    for sidecar in (
        answer.with_suffix(".component_map.json"),
        answer.with_suffix(".assumptions.json"),
        answer.with_suffix(".trainer.json"),
    ):
        if sidecar.is_file():
            shutil.copy2(sidecar, dest / sidecar.name)
    return copied_trainer, copied_answer


def _assert_readable_driver_layout(sheet) -> None:
    long_cells = 0
    for row in sheet.iter_rows(min_col=1, max_col=2, max_row=sheet.max_row or 1):
        for cell in row:
            value = cell.value
            if not isinstance(value, str) or value.startswith("="):
                continue
            if len(value) < 80:
                continue
            long_cells += 1
            assert cell.alignment.wrap_text is True, cell.coordinate
            height = sheet.row_dimensions[cell.row].height or 15
            assert height > 15, (cell.coordinate, len(value), height)
    assert long_cells >= 1


