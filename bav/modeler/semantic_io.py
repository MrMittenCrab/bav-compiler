"""Load runtime semantic component maps from workbooks and sidecars."""

from __future__ import annotations

from pathlib import Path

from bav.modeler.engine.semantic_map import ResolvedComponent, SemanticMap

_BAV_TRAINER_SUFFIX = "_BAV_Trainer"
_ANSWER_KEY_SUFFIX = "_Answer_Key"
_BAV_SUFFIX = "_BAV"
_TRAINER_SUFFIX = "_Trainer"


def supporting_dir(workbook_path: Path) -> Path:
    return Path(workbook_path).parent / "supporting"


def sidecar_paths(workbook_path: Path) -> tuple[Path, Path, Path]:
    """Return component_map, assumptions, and rowmap paths for a workbook."""
    workbook_path = Path(workbook_path)
    supporting = supporting_dir(workbook_path)
    supporting_map = supporting / "component_map.json"
    if supporting_map.is_file() or (supporting / "assumptions.json").is_file() or (
        supporting / "rowmap.json"
    ).is_file():
        return (
            supporting_map,
            supporting / "assumptions.json",
            supporting / "rowmap.json",
        )
    return (
        workbook_path.with_suffix(".component_map.json"),
        workbook_path.with_suffix(".assumptions.json"),
        workbook_path.parent / "rowmap.json",
    )


def component_map_path_for(workbook_path: Path) -> Path:
    """Sidecar path: supporting/component_map.json or foo.component_map.json."""
    return sidecar_paths(workbook_path)[0]


def company_stem_from_output(stem: str) -> str:
    """Strip product suffixes from a requested build stem."""
    for suffix in (
        _BAV_TRAINER_SUFFIX,
        _ANSWER_KEY_SUFFIX,
        _BAV_SUFFIX,
        _TRAINER_SUFFIX,
    ):
        if stem.endswith(suffix):
            return stem[: -len(suffix)]
    return stem


def resolve_pair_paths(output_path: Path) -> tuple[Path, Path]:
    """Resolve Trainer and BAV paths from a requested build output.

    Returns ``(trainer_path, bav_path)``. Ordinary and explicit builds use
    ``<Company>_BAV.xlsx`` and ``<Company>_BAV_Trainer.xlsx``.
    """
    output_path = Path(output_path)
    suffix = output_path.suffix or ".xlsx"
    company = company_stem_from_output(output_path.stem)
    parent = output_path.parent
    trainer_path = parent / f"{company}{_BAV_TRAINER_SUFFIX}{suffix}"
    bav_path = parent / f"{company}{_BAV_SUFFIX}{suffix}"
    return trainer_path, bav_path


def bav_path_for(training_path: Path) -> Path:
    """Infer matching BAV, falling back to a committed legacy Answer Key."""
    training_path = Path(training_path)
    _, bav_path = resolve_pair_paths(training_path)
    if bav_path.exists():
        return bav_path
    company = company_stem_from_output(training_path.stem)
    legacy = training_path.parent / f"{company}{_ANSWER_KEY_SUFFIX}{training_path.suffix or '.xlsx'}"
    if legacy.exists():
        return legacy
    return bav_path


def answer_key_path_for(training_path: Path) -> Path:
    """Backward-compatible alias — the BAV replaces the former Answer Key."""
    return bav_path_for(training_path)


def reference_workbook_path(training_path: Path) -> Path:
    """Backward-compatible alias for the matching professional BAV."""
    return bav_path_for(training_path)


def load_semantic_map(workbook_path: Path) -> SemanticMap:
    """Load component map from sidecar or embedded _ComponentMap sheet."""
    sidecar = component_map_path_for(workbook_path)
    if sidecar.exists():
        return SemanticMap.load_json(sidecar)
    return SemanticMap.from_workbook(workbook_path)


def get_component(workbook_path: Path, component_id: str) -> ResolvedComponent:
    return load_semantic_map(workbook_path).get(component_id)


def parse_cell_ref(cell_ref: str) -> tuple[int, int]:
    from openpyxl.utils import column_index_from_string

    col = "".join(c for c in cell_ref if c.isalpha())
    row = int("".join(c for c in cell_ref if c.isdigit()))
    return row, column_index_from_string(col)


def group_components_by_family(smap: SemanticMap) -> list[dict]:
    """Group concrete ResolvedComponents into conceptual schedule rows."""
    from bav.modeler.engine.component_catalog import (
        CAPEX_COMPONENT_CATALOG,
        COMPONENT_CATALOG,
        DEFERRED_TAX_COMPONENT_CATALOG,
        FIXED_ASSET_COMPONENT_CATALOG,
        GEOGRAPHIC_SEGMENT_COMPONENT_CATALOG,
        STORE_COUNT_COMPONENT_CATALOG,
        REVENUE_STORE_COMPONENT_CATALOG,
        COMPARABLE_SALES_COMPONENT_CATALOG,
        SALES_PER_SQUARE_FOOT_COMPONENT_CATALOG,
        REVENUE_DRIVER_COMPONENT_CATALOG,
        is_operating_kpi_source_identity,
        GOODWILL_INTANGIBLES_COMPONENT_CATALOG,
        LEASE_LIABILITY_COMPONENT_CATALOG,
        LEASE_REPAYMENT_COMPONENT_CATALOG,
        LEASE_ROU_COMPONENT_CATALOG,
        NORMALIZATION_COMPONENT_CATALOG,
        NORMALIZED_PER_SHARE_COMPONENT_CATALOG,
        OWNERSHIP_ATTRIBUTION_COMPONENT_CATALOG,
        PER_SHARE_ATTRIBUTION_COMPONENT_CATALOG,
        PER_SHARE_COMPONENT_CATALOG,
        PROFITABILITY_CHANGE_COMPONENT_CATALOG,
        PROFITABILITY_DRIVER_COMPONENT_CATALOG,
        QUALITY_CHANGE_COMPONENT_CATALOG,
        QUALITY_COMPONENT_CATALOG,
        ROE_ATTRIBUTION_COMPONENT_CATALOG,
        WORKING_CAPITAL_COMPONENT_CATALOG,
    )

    by_family: dict[str, list[ResolvedComponent]] = {}
    for comp in smap.all_ordered():
        by_family.setdefault(comp.family_id or comp.id, []).append(comp)

    family_meta = {f.id: f for f in COMPONENT_CATALOG}
    family_meta.update({f.id: f for f in NORMALIZATION_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in QUALITY_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in QUALITY_CHANGE_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in WORKING_CAPITAL_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in PROFITABILITY_DRIVER_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in PROFITABILITY_CHANGE_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in ROE_ATTRIBUTION_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in PER_SHARE_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in PER_SHARE_ATTRIBUTION_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in NORMALIZED_PER_SHARE_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in FIXED_ASSET_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in LEASE_LIABILITY_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in LEASE_ROU_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in OWNERSHIP_ATTRIBUTION_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in GOODWILL_INTANGIBLES_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in DEFERRED_TAX_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in CAPEX_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in LEASE_REPAYMENT_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in GEOGRAPHIC_SEGMENT_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in STORE_COUNT_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in REVENUE_STORE_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in COMPARABLE_SALES_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in SALES_PER_SQUARE_FOOT_COMPONENT_CATALOG})
    family_meta.update({f.id: f for f in REVENUE_DRIVER_COMPONENT_CATALOG})
    groups: list[dict] = []
    for family_id, comps in by_family.items():
        comps = [c for c in comps if not is_operating_kpi_source_identity(c)]
        if not comps:
            continue
        comps = sorted(comps, key=lambda c: (c.period_index is None, c.period_index or 0, c.order))
        first = comps[0]
        family = family_meta.get(family_id)
        ends = [c.period_end for c in comps if c.period_end]
        if ends:
            year_span = _period_scope_label(ends, len(comps))
        else:
            year_span = f"{len(comps)} cells"
        dep_ids: list[str] = []
        if family is not None:
            for dep in family.depends_on_current + family.depends_on_previous:
                if dep not in dep_ids:
                    dep_ids.append(dep)
        else:
            for dep in first.depends_on:
                fam = dep.split("__", 1)[0]
                if fam not in dep_ids:
                    dep_ids.append(fam)
        groups.append(
            {
                "family_id": family_id,
                "family_order": first.family_order or first.order,
                "title": first.title,
                "period_scope": year_span,
                "tab": first.tab,
                "practice_cells": _format_practice_cells(comps),
                "depends_on": ", ".join(dep_ids) if dep_ids else "—",
                "count": len(comps),
                "components": comps,
            }
        )
    groups.sort(key=lambda g: g["family_order"])
    return groups


def _period_scope_label(period_ends: list[str], count: int) -> str:
    years = []
    for end in period_ends:
        years.append(end[:4] if len(end) >= 4 else end)
    if len(years) == 1:
        return f"{years[0]} ({count} cells)"
    return f"{years[0]}–{years[-1]} ({count} cells)"


def _format_practice_cells(comps: list[ResolvedComponent]) -> str:
    from openpyxl.utils import column_index_from_string, get_column_letter

    if not comps:
        return "—"
    cells = [c.cell for c in comps]
    if len(cells) == 1:
        return cells[0]

    parsed = []
    for cell in cells:
        col = "".join(ch for ch in cell if ch.isalpha())
        row = int("".join(ch for ch in cell if ch.isdigit()))
        parsed.append((row, column_index_from_string(col), cell))

    rows = {p[0] for p in parsed}
    if len(rows) == 1:
        cols = sorted(p[1] for p in parsed)
        if cols == list(range(cols[0], cols[0] + len(cols))):
            row = next(iter(rows))
            start = f"{get_column_letter(cols[0])}{row}"
            end = f"{get_column_letter(cols[-1])}{row}"
            return f"{start}:{end}"
    return ", ".join(cells)
