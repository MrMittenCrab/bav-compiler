"""Professional BAV Overview opening presentation."""

from __future__ import annotations

from openpyxl.styles import Alignment

BAV_OPENING_SHEET = "Overview"
_HIDDEN_PREFIX = "_"
_SELECTED_FALLBACK_SCHEDULES = (
    "Income Statement",
    "Balance Sheet",
    "Cash Flow Statement",
    "Condensed Financials",
    "ALT DuPont",
    "Geographic Segment Analysis",
    "Store Count Analysis",
    "Revenue Driver Analysis",
)


def _add_bav_opening(wb, financials) -> None:
    """Professional cover: identity, source-backed interpretation, navigation."""
    for name in (BAV_OPENING_SHEET, "Trainer"):
        if name in wb.sheetnames:
            del wb[name]
    fin = financials
    periods = list(fin.periods or [])
    labels = []
    for period in periods:
        label = (getattr(period, "label", None) or "").strip()
        end = getattr(period, "end_date", None)
        if label:
            labels.append(label)
        elif end is not None:
            labels.append(end.isoformat() if hasattr(end, "isoformat") else str(end))
    if labels:
        coverage = (
            labels[0]
            if len(labels) == 1
            else f"{labels[0]} – {labels[-1]} ({len(labels)} periods)"
        )
    else:
        coverage = "No admitted historical periods"
    company = (fin.company_name or fin.ticker or "Company").strip()
    ticker = (fin.ticker or "").strip()
    identity = f"{company} ({ticker})" if ticker and ticker.casefold() != company.casefold() else company
    currency = (fin.currency or "").strip() or "unspecified currency"
    units = (fin.units or "").strip() or "unspecified units"
    structure = [
        title
        for title in wb.sheetnames
        if not title.startswith(_HIDDEN_PREFIX)
        and title != "Build Status"
        and wb[title].sheet_state == "visible"
    ]
    from composer.overview import PROFESSIONAL_FALLBACK
    from director.driver_assessment import (
        complete_historical_strategy_synthesis,
        complete_revenue_driver_analysis,
    )
    from modeler.revenue_driver import strategy_synthesis_applicable

    ws = wb.create_sheet(BAV_OPENING_SHEET, 0)
    ws.sheet_view.showGridLines = False
    wrap = Alignment(wrap_text=True, vertical="top")
    label_width = 28.0
    narrative_width = 88.0
    ws.column_dimensions["A"].width = label_width
    ws.column_dimensions["B"].width = narrative_width

    def _wrapped_height(text: str, *, width: float) -> float:
        chars_per_line = max(24, int(width * 0.9))
        paragraphs = str(text or "").splitlines() or [""]
        lines = 0
        for paragraph in paragraphs:
            length = max(len(paragraph), 1)
            lines += max(1, (length + chars_per_line - 1) // chars_per_line)
        return max(18.0, 15.0 * lines + 8.0)

    def _raise_row(row: int, height: float) -> None:
        current = ws.row_dimensions[row].height
        ws.row_dimensions[row].height = max(float(current or 0), height)

    def _write(row: int, column: int, text: str, *, width: float) -> None:
        cell = ws.cell(row=row, column=column, value=text)
        cell.alignment = wrap
        if text:
            _raise_row(row, _wrapped_height(text, width=width))

    def _label(row: int, text: str) -> None:
        _write(row, 1, text, width=label_width)

    def _narrative(row: int, text: str) -> None:
        _write(row, 2, text, width=narrative_width)
        label = ws.cell(row=row, column=1).value
        if isinstance(label, str) and label:
            _raise_row(row, _wrapped_height(label, width=label_width))

    def _wide(row: int, text: str) -> None:
        cell = ws.cell(row=row, column=1, value=text)
        cell.alignment = wrap
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        _raise_row(row, _wrapped_height(text, width=label_width + narrative_width))

    _wide(1, "BAV")
    _wide(2, identity)
    _wide(3, f"Historical coverage: {coverage}")
    _wide(4, f"Units: {currency}; {units}")

    cursor = 6
    navigation: list[str] = []
    if strategy_synthesis_applicable(fin):
        analysis = complete_revenue_driver_analysis(fin)
        synthesis = complete_historical_strategy_synthesis(fin, analysis)
        _label(cursor, "Historical reading")
        _narrative(cursor, synthesis.lead)
        cursor += 2
        limits = " ".join(
            part
            for part in (
                synthesis.limits,
                synthesis.productivity_gap,
                synthesis.untested,
            )
            if part
        )
        _label(cursor, "Evidence limits")
        _narrative(cursor, limits)
        cursor += 2
        navigation = [
            name for name in synthesis.navigation if name in structure
        ]
    else:
        _label(cursor, "Historical reading")
        _narrative(cursor, PROFESSIONAL_FALLBACK)
        cursor += 2

    if not navigation:
        navigation = [
            name for name in _SELECTED_FALLBACK_SCHEDULES if name in structure
        ]
    if not navigation:
        navigation = list(structure)[:8]
    _wide(cursor, "Supporting schedules")
    cursor += 1
    for name in navigation:
        cell = ws.cell(row=cursor, column=1, value=name)
        cell.alignment = wrap
        if name in wb.sheetnames:
            cell.hyperlink = f"#'{name}'!A1"
        cursor += 1
    cell = ws.cell(row=cursor, column=1, value="Build Status")
    cell.alignment = wrap
    cell.hyperlink = "#'Build Status'!A1"
