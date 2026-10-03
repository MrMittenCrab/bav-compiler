"""Disclosure-led historical revenue-driver tests and workbook presentation."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path
import shutil

import pytest
from openpyxl import load_workbook

from extractor.data.historical_strategy import (
    ROLE_OBJECTIVE,
    ROLE_OPERATING_USE,
    ROLE_STRATEGY,
    THEME_COMPARABLE_SALES,
    THEME_GEOGRAPHIC_GROWTH,
    THEME_OPERATING_MARGIN,
    THEME_PRODUCTIVITY,
    THEME_STORE_EXPANSION,
    HistoricalStrategyData,
    HistoricalStrategyDisclosure,
    disclosure_locator,
    load_strategy_disclosures,
)
from modeler.data.standardized_io import standardized_from_payload, standardized_to_payload
from modeler.data.interface import (
    HistoricalManagementKpiDeferredDisagreement,
    HistoricalManagementKpiDeferredMember,
)
from modeler.engine.component_catalog import (
    REVENUE_DRIVER_COMPSALES_DIFFERENCE_FAMILY_ID,
    REVENUE_DRIVER_COMPSALES_FAMILY_ID,
    REVENUE_DRIVER_COMPONENT_CATALOG,
    REVENUE_DRIVER_GEO_CONTRIBUTION_FAMILY_ID,
    REVENUE_DRIVER_REVENUE_GROWTH_FAMILY_ID,
    REVENUE_DRIVER_RPS_FAMILY_ID,
    REVENUE_DRIVER_SHEET_NAME,
    REVENUE_DRIVER_STORE_DIFFERENCE_FAMILY_ID,
    REVENUE_DRIVER_STORE_GROWTH_FAMILY_ID,
    SemanticCellRef,
    expand_revenue_driver_specs,
    is_operating_kpi_source_identity,
    resolve_revenue_driver_link_formula,
    revenue_driver_component_id,
)
from modeler.workbook import ReferenceModelBuilder
from modeler.ingestion.management_kpi_identity import (
    FAMILY_COMPARABLE_SALES_GROWTH,
    FAMILY_SALES_PER_SQUARE_FOOT,
    POP_COMPANY_OPERATED_STORES,
    POP_STORES_AND_DTC,
    POP_STORES_AND_ECOMMERCE,
)
from modeler.line_resolver import MissingLineError
from modeler.period_axis import PeriodAxisError
from composer.overview import (
    DEFERRED_SPSF_LINK,
    PROFESSIONAL_FALLBACK,
    UNTESTED_INITIATIVES,
    WHAT_HISTORY_ESTABLISHES,
)
from composer.revenue_driver import (
    ADDITIONAL_SPSF,
    FAILED_COMPSALES,
    FAILED_SPSF_GROWTH,
    FAILED_STORE_GROWTH,
    SCOPE_NOTE,
)
from director.driver_assessment import (
    complete_historical_strategy_synthesis as compute_historical_strategy_synthesis,
    complete_revenue_driver_analysis as compute_revenue_driver_analysis,
)
from interpreter.revenue_driver import (
    VERDICT_CONTRADICTED,
    VERDICT_INSUFFICIENT,
    VERDICT_MIXED,
    VERDICT_SUPPORTED,
)
from modeler.revenue_driver import revenue_driver_applicable, strategy_synthesis_applicable
from modeler.tests.test_capex import P1, P2, _tiny
from modeler.tests.test_historical_segment import _corp_values, _snapshot
from composer.tests.test_learner_ready_presentation import (
    _assert_answer_key_no_yellow,
    assert_bav_has_no_exercise_framing,
)
from modeler.tests.test_geographic_segment_workbook import _geo_tiny
from modeler.tests.test_operating_kpi_facts import _kpi_model_observation
from modeler.tests.test_operating_kpi_management_history import _compsales, _spsf
from modeler.tests.test_operating_kpi_relationships import _fin_with_relationship
from legacy.trainer.checker import check_workbook
from modeler.tests.test_normalization import _inject_formula_and_cached_value
from modeler.semantic_io import load_semantic_map, parse_cell_ref
from legacy.trainer.derive import build_training_workbook, derive_trainer_workbook

ROOT = Path(__file__).resolve().parents[2]
FR_JSON = ROOT / "build" / "input" / "fast_retailing" / "reconciled" / "standardized.json"
LULU_JSON = ROOT / "core" / "tests" / "fixtures" / "ordinary_reconcile" / "lululemon" / "standardized.json"
LULU_DISCLOSURES = (
    ROOT / "core" / "tests" / "fixtures" / "strategy" / "lululemon_management_disclosures.json"
)
P0 = date(2023, 12, 31)


def _disclosure(
    theme: str,
    *,
    role: str = ROLE_STRATEGY,
    text: str = "We open stores.",
    period: date = P2,
) -> HistoricalStrategyDisclosure:
    return HistoricalStrategyDisclosure(
        theme=theme,
        role=role,
        text=text,
        period=period,
        source_file="example.pdf",
        page_reference="Form 10-K p. 1",
        section="Item 1",
    )


def _with_strategy(fin, *disclosures: HistoricalStrategyDisclosure):
    fin.historical_strategy = HistoricalStrategyData(disclosures=disclosures)
    return fin


def _store_fin(
    counts: dict[date, float],
    revenue: dict[date, float | None],
    *disclosures: HistoricalStrategyDisclosure,
    management=None,
):
    observations = tuple(
        _kpi_model_observation(period, value) for period, value in counts.items()
    )
    fin = _fin_with_relationship(
        *observations,
        revenue=revenue,
        extra_periods=list(counts),
        management=management,
    )
    return _with_strategy(fin, *disclosures)


def _deferred_member(
    *,
    locator: str,
    definition_text: str,
    presentation_role: str = "current",
    extraction_document: str = "example.json",
    page_reference: str = "Form 10-K p. 10",
    physical_page_mapping: str = "10",
    reported_value: float | None = 1580.0,
) -> HistoricalManagementKpiDeferredMember:
    return HistoricalManagementKpiDeferredMember(
        locator=locator,
        extraction_document=extraction_document,
        page_reference=page_reference,
        physical_page_mapping=physical_page_mapping,
        presentation_role=presentation_role,
        definition_text=definition_text,
        population=POP_COMPANY_OPERATED_STORES,
        unit="USD_per_square_foot",
        basis="reported",
        calendar_week_adjustment="included",
        calendar_reporting_basis="52_week",
        reported_value=reported_value,
    )


def _deferred_disagreement(
    *,
    family: str = FAMILY_SALES_PER_SQUARE_FOOT,
    period: date = P2,
    reasons: tuple[str, ...] = ("definition_mismatch", "ordinary_disagreement"),
    members: tuple[HistoricalManagementKpiDeferredMember, ...] | None = None,
) -> HistoricalManagementKpiDeferredDisagreement:
    return HistoricalManagementKpiDeferredDisagreement(
        family=family,
        period=period,
        reasons=list(reasons),
        members=list(
            members
            or (
                _deferred_member(
                    locator=f"{family}:current:{period.isoformat()}",
                    definition_text="average during the year",
                    presentation_role="current",
                ),
                _deferred_member(
                    locator=f"{family}:prior:{period.isoformat()}",
                    definition_text="average ending square footage",
                    presentation_role="prior",
                    page_reference="Form 10-K p. 11",
                    physical_page_mapping="11",
                ),
            )
        ),
    )


def _attach_deferred(fin, *items: HistoricalManagementKpiDeferredDisagreement):
    data = fin.historical_operating_kpis
    assert data is not None
    data.deferred_disagreements = list(items)
    return fin


def test_opening_without_strategy_stays_professional_fallback(tmp_path):
    fin = _tiny()
    assert not strategy_synthesis_applicable(fin)
    trainer, answer = build_training_workbook(fin, tmp_path / "NO_STRATEGY.xlsx")
    awb = load_workbook(answer, data_only=False)
    opening = " ".join(
        str(cell.value or "") for row in awb["Overview"].iter_rows() for cell in row
    )
    assert awb["Overview"]["A1"].value == "BAV"
    assert PROFESSIONAL_FALLBACK in opening
    assert "Historical reading" in opening
    assert "HISTORICAL REVENUE AND DISCLOSED STRATEGY" not in opening
    assert "Schedules: " not in opening
    assert "Management statement" not in opening
    assert "Historical finding" not in opening
    assert any(
        cell.hyperlink
        and "Build Status" in str(getattr(cell.hyperlink, "target", "") or cell.hyperlink)
        for row in awb["Overview"].iter_rows()
        for cell in row
        if cell.hyperlink
    )
    awb.close()
    derive_trainer_workbook(answer, trainer)


def test_strategy_synthesis_connects_findings_without_claiming_outcomes():
    from modeler.tests.test_geographic_segment_workbook import _geo_tiny

    fin = _geo_tiny(
        _snapshot(P1, values=_corp_values(rev=(80.0, 25.0, 15.0))),
        _snapshot(P2, values=_corp_values(rev=(110.0, 15.0, 15.0))),
    )
    store = _store_fin(
        {P1: 10, P2: 12},
        {P1: 120.0, P2: 140.0},
        _disclosure(
            THEME_STORE_EXPANSION,
            role=ROLE_OBJECTIVE,
            text="We plan to open stores.",
        ),
        management=[
            _compsales(period=P1, value=4.0, geography="global", basis="reported"),
            _compsales(period=P2, value=5.0, geography="global", basis="reported"),
        ],
    )
    fin.historical_operating_kpis = store.historical_operating_kpis
    store_statement = _disclosure(
        THEME_STORE_EXPANSION,
        role=ROLE_OBJECTIVE,
        text="We plan to open stores.",
    )
    compsales_statement = _disclosure(
        THEME_COMPARABLE_SALES, role=ROLE_OPERATING_USE
    )
    productivity_statement = _disclosure(
        THEME_PRODUCTIVITY, role=ROLE_OPERATING_USE
    )
    geo_statement = _disclosure(THEME_GEOGRAPHIC_GROWTH)
    fin.historical_strategy = HistoricalStrategyData(
        disclosures=(
            store_statement,
            compsales_statement,
            productivity_statement,
            geo_statement,
        )
    )
    analysis = compute_revenue_driver_analysis(fin)
    synthesis = compute_historical_strategy_synthesis(fin, analysis)
    assert strategy_synthesis_applicable(fin)
    assert "historical growth pattern" in synthesis.lead
    assert "do not establish causal drivers" in synthesis.lead
    assert "comprehensive strategy execution" in synthesis.lead
    assert synthesis.untested == UNTESTED_INITIATIVES
    assert synthesis.limits == WHAT_HISTORY_ESTABLISHES
    by_theme = {item.theme: item for item in synthesis.interpretations}
    assert set(by_theme) == {
        THEME_STORE_EXPANSION,
        THEME_COMPARABLE_SALES,
        THEME_PRODUCTIVITY,
        THEME_GEOGRAPHIC_GROWTH,
    }
    store_row = by_theme[THEME_STORE_EXPANSION]
    assert store_statement.text in store_row.management_statement
    assert disclosure_locator(store_statement) in store_row.management_statement
    assert analysis.tests[0].finding in store_row.finding
    assert REVENUE_DRIVER_SHEET_NAME in store_row.finding
    assert "not achieved historical outcomes" in store_row.inference
    assert "not new-store contribution" in store_row.inference
    geo_row = by_theme[THEME_GEOGRAPHIC_GROWTH]
    assert "contributed negatively" in geo_row.inference
    assert "arithmetic decomposition" in geo_row.inference
    productivity = by_theme[THEME_PRODUCTIVITY]
    assert "cannot test" in productivity.inference
    assert "not independent productivity evidence" in synthesis.productivity_gap
    assert DEFERRED_SPSF_LINK not in synthesis.productivity_gap
    assert "average during the year" not in synthesis.productivity_gap
    assert "ordinary_disagreement" not in synthesis.productivity_gap


def _synthesis_record(synthesis):
    return {
        "lead": synthesis.lead,
        "productivity_gap": synthesis.productivity_gap,
        "untested": synthesis.untested,
        "limits": synthesis.limits,
        "navigation": synthesis.navigation,
        "interpretations": tuple(
            (
                item.theme,
                item.heading,
                item.management_statement,
                item.finding,
                item.inference,
                item.supporting_schedules,
                item.counterexample,
            )
            for item in synthesis.interpretations
        ),
    }


def test_composer_renders_supplied_judgment_without_interpretation_or_compute(
    monkeypatch,
):
    from dataclasses import replace

    from composer.overview import compute_historical_strategy_synthesis as compose
    from director.driver_assessment import complete_revenue_driver_analysis
    from interpreter.historical_strategy import interpret_historical_strategy

    fin = _store_fin(
        {P0: 10, P1: 12, P2: 15},
        {P0: 100.0, P1: 130.0, P2: 160.0},
        _disclosure(THEME_STORE_EXPANSION),
    )
    analysis = complete_revenue_driver_analysis(fin)
    judgment = interpret_historical_strategy(analysis.tests)

    def boom(*args, **kwargs):
        raise AssertionError("Composer must not interpret or compute")

    monkeypatch.setattr(
        "interpreter.historical_strategy.interpret_historical_strategy", boom
    )
    monkeypatch.setattr(
        "modeler.revenue_driver.compute_revenue_driver_analysis", boom
    )
    monkeypatch.setattr(
        "director.driver_assessment.complete_revenue_driver_analysis", boom
    )
    monkeypatch.setattr(
        "director.driver_assessment.interpret_historical_strategy", boom
    )

    synthesis = compose(fin, analysis, judgment)
    assert "historical growth pattern" in synthesis.lead
    assert synthesis.untested == UNTESTED_INITIATIVES
    assert synthesis.limits == WHAT_HISTORY_ESTABLISHES

    overview_src = (ROOT / "composer" / "overview.py").read_text()
    assert "interpret_historical_strategy" not in overview_src
    assert "complete_revenue_driver" not in overview_src
    assert "compute_revenue_driver_analysis" not in overview_src

    tiny = _tiny()
    with pytest.raises(ValueError, match="admitted strategy disclosures"):
        compose(tiny)
    with pytest.raises(ValueError, match="admitted strategy disclosures"):
        compose(tiny, analysis, judgment)
    with pytest.raises(ValueError, match="completed revenue-driver analysis"):
        compose(fin)
    empty = replace(analysis, tests=())
    with pytest.raises(ValueError, match="at least one driver test"):
        compose(fin, empty)
    with pytest.raises(ValueError, match="completed historical-strategy judgment"):
        compose(fin, analysis)


DRIVER_FAMILY_IDS = (
    REVENUE_DRIVER_STORE_GROWTH_FAMILY_ID,
    REVENUE_DRIVER_REVENUE_GROWTH_FAMILY_ID,
    REVENUE_DRIVER_STORE_DIFFERENCE_FAMILY_ID,
    REVENUE_DRIVER_COMPSALES_FAMILY_ID,
    REVENUE_DRIVER_COMPSALES_DIFFERENCE_FAMILY_ID,
    REVENUE_DRIVER_RPS_FAMILY_ID,
    REVENUE_DRIVER_GEO_CONTRIBUTION_FAMILY_ID,
)


def _copy_workbooks(trainer: Path, answer: Path, dest: Path) -> tuple[Path, Path]:
    dest.mkdir()
    copied_trainer = dest / trainer.name
    copied_answer = dest / answer.name
    shutil.copy2(trainer, copied_trainer)
    shutil.copy2(answer, copied_answer)
    for sidecar in (
        answer.with_suffix(".component_map.json"),
        answer.with_suffix(".assumptions.json"),
        answer.with_suffix(".trainer.json"),
    ):
        if sidecar.is_file():
            shutil.copy2(sidecar, dest / sidecar.name)
    return copied_trainer, copied_answer


def _assert_readable_driver_layout(sheet) -> None:
    long_cells = 0
    for row in sheet.iter_rows(min_col=1, max_col=2, max_row=sheet.max_row or 1):
        for cell in row:
            value = cell.value
            if not isinstance(value, str) or value.startswith("="):
                continue
            if len(value) < 80:
                continue
            long_cells += 1
            assert cell.alignment.wrap_text is True, cell.coordinate
            height = sheet.row_dimensions[cell.row].height or 15
            assert height > 15, (cell.coordinate, len(value), height)
    assert long_cells >= 1


def test_phase_boundaries_and_structured_flags_survive_wording():
    import ast
    from pathlib import Path

    from composer.revenue_driver import word_identity_assessment, word_revenue_assessment
    from interpreter.revenue_driver import interpret_revenue_assessment, interpret_theme
    from modeler.revenue_driver import (
        _identity_assessment,
        compute_revenue_driver_analysis as compute_numeric,
    )

    root = Path(__file__).resolve().parents[2]
    modeler_src = ast.parse((root / "modeler" / "revenue_driver.py").read_text())
    interpreter_src = ast.parse((root / "interpreter" / "revenue_driver.py").read_text())
    modeler_imports = {
        alias.name
        for node in modeler_src.body
        if isinstance(node, ast.ImportFrom) and node.module
        for alias in node.names
    }
    modeler_modules = {
        node.module
        for node in modeler_src.body
        if isinstance(node, ast.ImportFrom) and node.module
    }
    interpreter_modules = {
        node.module
        for node in interpreter_src.body
        if isinstance(node, ast.ImportFrom) and node.module
    }
    assert not any(module.startswith("interpreter") or module.startswith("composer") for module in modeler_modules)
    assert not any(module.startswith("composer") for module in interpreter_modules)
    assert "THEME_LABELS" not in modeler_imports

    fin = _store_fin(
        {P0: 10, P1: 12, P2: 15},
        {P0: 100.0, P1: 130.0, P2: 160.0},
        _disclosure(THEME_STORE_EXPANSION),
    )
    numeric = compute_numeric(fin)
    bundle = numeric.theme_observations[0]
    interpretation = interpret_theme(bundle)
    assert any(item.counterexample for item in interpretation.observations.observations)
    identity = _identity_assessment(
        bundle.theme,
        numeric.geographic_reconstruction,
        numeric.footprint_identity,
    )
    from composer.revenue_driver import word_theme

    test = word_theme(interpretation)
    if identity is not None:
        assessment = word_identity_assessment(identity, test)
        assert assessment.kind == identity.kind
        assert assessment.established == identity.established
    else:
        decision = interpret_revenue_assessment(interpretation)
        assessment = word_revenue_assessment(decision, test)
        assert assessment.kind == decision.kind
        assert assessment.established == decision.established
    assert any(item.note and item.counterexample == item.counterexample for item in test.observations)

    deferred = _deferred_disagreement(period=P1)
    deferred_fin = _attach_deferred(
        _store_fin(
            {P0: 10, P1: 12, P2: 15},
            {P0: 100.0, P1: 130.0, P2: 160.0},
            _disclosure(THEME_PRODUCTIVITY),
            management=[_spsf(period=P2, value=1500)],
        ),
        deferred,
    )
    deferred_numeric = compute_numeric(deferred_fin)
    productivity = interpret_theme(deferred_numeric.theme_observations[0])
    assert productivity.has_deferred_spsf is True
    from composer.revenue_driver import word_theme as word_prod

    worded = word_prod(productivity)
    assert worded.has_deferred_spsf is True
    assert any(item.startswith("Deferred ") for item in worded.limitations)
    interpreter_text = (root / "interpreter" / "revenue_driver.py").read_text()
    historical_text = (root / "interpreter" / "historical_strategy.py").read_text()
    overview_text = (root / "composer" / "overview.py").read_text()
    assert "_is_counterexample_note" not in interpreter_text
    assert '"Deferred' not in interpreter_text
    assert '"Deferred' not in historical_text
    assert "THEME_LABELS" not in historical_text
    assert "composer" not in historical_text
    assert "interpret_historical_strategy" not in overview_text
    checker_src = (root / "core" / "trainer" / "checker.py").read_text()
    expected_src = (root / "core" / "model" / "historical_expected.py").read_text()
    catalog_src = (root / "core" / "engine" / "component_catalog.py").read_text()
    reference_src = (root / "core" / "engine" / "reference_model.py").read_text()
    assert "word_lead" not in checker_src
    assert "word_verdict_inference" not in expected_src
    assert "interpret_historical_strategy" not in catalog_src
    assert "word_verdict_inference" not in reference_src

    from modeler.revenue_driver import _conflicting_qualifier_fields

    qualifier_item = _deferred_disagreement(
        members=(
            _deferred_member(
                locator="a",
                definition_text="same definition",
                presentation_role="current",
            ),
            _deferred_member(
                locator="b",
                definition_text="same definition",
                presentation_role="prior",
            ),
        )
    )
    qualifier_item.members[1].population = POP_STORES_AND_DTC
    assert _conflicting_qualifier_fields(qualifier_item) == ("population",)

    geo_fin = _geo_tiny(
        _snapshot(P1, values=_corp_values(rev=(100.0, 20.0, 10.0))),
        _snapshot(P2, values=_corp_values(rev=(90.0, 25.0, 10.0))),
    )
    _with_strategy(geo_fin, _disclosure(THEME_GEOGRAPHIC_GROWTH))
    geo_numeric = compute_numeric(geo_fin)
    geo_interp = interpret_theme(geo_numeric.theme_observations[0])
    assert any(
        item.consistent is False and item.counterexample
        for item in geo_interp.observations.observations
    )
    geo_test = word_theme(geo_interp)
    assert "did not grow" in geo_test.finding
    geo_identity = _identity_assessment(
        geo_numeric.theme_observations[0].theme,
        geo_numeric.geographic_reconstruction,
        geo_numeric.footprint_identity,
    )
    geo_assessment = word_identity_assessment(geo_identity, geo_test)
    assert geo_assessment.kind == geo_identity.kind
    assert geo_assessment.established == geo_identity.established
    assert geo_assessment.contradictions == geo_test.finding

    prod_fin = _store_fin(
        {P1: 10, P2: 15},
        {P1: 100.0, P2: 110.0},
        _disclosure(THEME_PRODUCTIVITY, role=ROLE_OPERATING_USE),
    )
    prod_complete = compute_revenue_driver_analysis(prod_fin).tests[0]
    assert prod_complete.rps_declines > 0
    assert prod_complete.assessment is not None
    assert prod_complete.assessment.contradictions == prod_complete.finding
    assert "declined in" in prod_complete.finding


