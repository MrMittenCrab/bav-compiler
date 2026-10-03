"""Step 9M.2.4.1.1.1.26 — source-supported reported operating-margin bridge."""

from __future__ import annotations

import copy
import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from director.current_build import prepare_company_input, resolve_company
from extractor.data.interface import (
    DocumentManifest,
    DocumentType,
)
from modeler.data.interface import (
    LineItem,
    StandardizedFinancials,
)
from modeler.data.standardized_io import standardized_from_payload, standardized_to_payload
from modeler.engine.component_catalog import (
    REPORTED_MARGIN_COMPONENT_CATALOG,
    expand_reported_margin_specs,
)
from modeler.workbook import ReferenceModelBuilder
from legacy.ingestion.manual_hk import HKManualDocumentAdapter
from modeler.historical_expected import reported_margin_expected_series
from modeler.line_resolver import AmbiguousLineError, MissingLineError, resolve_line
from modeler.period_axis import canonical_fiscal_periods
from modeler.ratio_values import UNDEFINED_RATIO
from composer.research.drivers import render_drivers_markdown
from director.driver_assessment import complete_reported_margin_series as compute_reported_margin_series
from director.research import complete_drivers_view as assemble_drivers_view
from modeler.reported_margin import (
    EXPENSE_PRESENTATION_POSITIVE,
    EXPENSE_PRESENTATION_SIGNED,
    GROSS_PROFIT_CONCEPT,
    OPERATING_PROFIT_CONCEPT,
    REVENUE_CONCEPT,
    analytical_expense,
    expense_presentation_factor,
    gross_margin_applicable,
    net_operating_expense_burden_applicable,
    publication_reconstruction_allowed,
    reported_margin_applicable,
    reported_margin_availability,
    reported_operating_margin_applicable,
    residual_blocks_reconstruction_claim,
    resolve_reported_margin_sources,
)
from modeler.source_values import MissingHistoricalValueError
from modeler.tests.test_capex import P1, P2, _dupont_row_by_label, _tiny
from modeler.tests.test_normalization import _inject_formula_and_cached_value
from legacy.trainer.checker import check_workbook
from modeler.semantic_io import load_semantic_map, parse_cell_ref
from modeler.semantic_io import group_components_by_family
from legacy.trainer.derive import build_training_workbook

ROOT = Path(__file__).resolve().parents[2]
FR_JSON = ROOT / "build" / "input" / "fast_retailing" / "reconciled" / "standardized.json"
DEMO_JSON = ROOT / "legacy" / "example" / "DEMO_HK_Standardized.json"


def _li(label, values, concept=""):
    return LineItem(label=label, values=values, concept=concept)


def _vals(a, b, *, missing_period=False, none_period=False):
    if missing_period:
        return {P1: a}
    if none_period:
        return {P1: a, P2: None}
    return {P1: a, P2: b}


def _add_margins(
    fin: StandardizedFinancials,
    *,
    revenue=(1000.0, 1100.0),
    gross=(560.0, 600.0),
    operating=(200.0, 180.0),
    revenue_concept=REVENUE_CONCEPT,
    gross_concept=GROSS_PROFIT_CONCEPT,
    operating_concept=OPERATING_PROFIT_CONCEPT,
    include_revenue=True,
    include_gross=True,
    include_operating=True,
    duplicate=None,
    missing_period=None,
    none_period=None,
    on_balance_sheet=None,
):
    rows = (
        ("revenue", include_revenue, "Revenue", revenue_concept, revenue),
        ("gross", include_gross, "Gross profit", gross_concept, gross),
        (
            "operating",
            include_operating,
            "Operating profit",
            operating_concept,
            operating,
        ),
    )
    stored = {}
    for key, include, label, concept, values in rows:
        if not include:
            continue
        item = _li(
            label,
            _vals(
                values[0],
                values[1],
                missing_period=missing_period == key,
                none_period=none_period == key,
            ),
            concept=concept,
        )
        if key == "revenue":
            existing = next(
                (
                    row
                    for row in fin.income_statement
                    if row.concept == "revenue" or row.label.lower() == "revenue"
                ),
                None,
            )
            if existing is not None:
                fin.income_statement.remove(existing)
        target = (
            fin.balance_sheet if on_balance_sheet == key else fin.income_statement
        )
        target.append(item)
        stored[key] = item
        if duplicate == key and on_balance_sheet != key:
            target.append(
                _li(
                    f"{label} duplicate",
                    _vals(values[0], values[1]),
                    concept=concept,
                )
            )
    return stored


LULU_RECONCILED = (
    ROOT / "build" / "input" / "lululemon" / "reconciled" / "standardized.json"
)


def _add_component_lines(
    fin: StandardizedFinancials,
    *,
    sga=(300.0, 300.0),
    impairment=(0.0, 50.0),
    other=(10.0, 0.0),
    omit_impairment=False,
    impairment_none_current=False,
):
    fin.income_statement.append(
        _li(
            "SG&A",
            _vals(sga[0], sga[1]),
            concept="selling_general_and_administrative_expenses",
        )
    )
    if not omit_impairment:
        fin.income_statement.append(
            _li(
                "Impairment",
                {P1: impairment[0], P2: None}
                if impairment_none_current
                else _vals(impairment[0], impairment[1]),
                concept="impairment_and_restructuring",
            )
        )
    fin.income_statement.append(
        _li(
            "Amortization",
            _vals(other[0], other[1]),
            concept="amortization_of_intangible_assets",
        )
    )


def test_rendered_component_contribution_schedule(tmp_path):
    company = resolve_company("Lululemon")
    fin = prepare_company_input(company, tmp_path / "input")
    trainer, answer = build_training_workbook(fin, tmp_path / "LULU_CONTRIB.xlsx")
    awb = load_workbook(answer, data_only=False)
    ws = awb["ALT DuPont"]
    assert any(
        ws.cell(row, 1).value == "COMPONENT OPERATING-MARGIN CONTRIBUTIONS"
        for row in range(1, (ws.max_row or 1) + 1)
    )
    gm_row = _dupont_row_by_label(ws, "Δ gross margin contribution")
    sga_row = _dupont_row_by_label(ws, "−Δ SG&A/revenue contribution")
    imp_row = _dupont_row_by_label(ws, "−Δ impairment/revenue contribution")
    other_row = _dupont_row_by_label(
        ws, "−Δ other operating items/revenue contribution"
    )
    sum_row = _dupont_row_by_label(ws, "Reconstructed contribution sum")
    resid_row = _dupont_row_by_label(
        ws, "Contribution residual (reported − reconstructed)"
    )
    bps_row = _dupont_row_by_label(ws, "Δ gross margin contribution (bps)")
    assert ws.cell(gm_row, 2).value in (None, "")
    gm_f = str(ws.cell(gm_row, 3).value)
    sga_f = str(ws.cell(sga_row, 3).value)
    assert gm_f.startswith("=")
    assert sga_f.startswith("=")
    assert "-(" in sga_f or sga_f.startswith("=-")
    assert "*10000" in str(ws.cell(bps_row, 3).value)
    assert '""' in str(ws.cell(sum_row, 3).value)
    assert "reported" not in str(ws.cell(resid_row, 3).value).lower() or "-" in str(
        ws.cell(resid_row, 3).value
    )
    awb.close()

    text = render_drivers_markdown(assemble_drivers_view(fin, company.name))
    main, appendix = text.split("## Appendix", 1)
    assert "Δgross margin" in appendix
    assert "−Δ(SG&A/revenue)" in appendix
    assert "−Δ(impairment/revenue)" in appendix
    assert "Reconstructed sum" in appendix
    assert "unrounded" in appendix
    assert "10,000" in appendix
    names = [
        line.split("|")[1].strip()
        for line in appendix.splitlines()
        if line.startswith("| ") and " | " in line
    ]
    assert names.count("component operating-margin identity") == 1
    assert names.count("latest adjacent operating-margin movement") == 1
    assert "component operating-margin contributions" in appendix
    assert "component operating-margin identity" not in main
    assert "## Kind" not in main and "### Residual" not in main


